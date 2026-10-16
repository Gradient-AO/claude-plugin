#!/usr/bin/env python3
"""Render, validate and create a visual contact sheet for a Gradient report."""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from report_layout import (
    LayoutPolicyError,
    build_post_render_layout_contract,
    cover_allowed_text,
    cover_fact_count_errors,
    cover_page_failures,
    lint_report_design,
    reader_text_failures,
    load_layout_budget,
    prepare_report_for_render,
    repair_preferred_breaks,
    resolve_layout_budget,
    validate_report_budget,
)
from report_theme import FONT_FAMILY, THEME_VERSION, load_branding

MAX_REPAIR_PASSES = 3

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def run(command: list[str]) -> subprocess.CompletedProcess[str]:
    """Run one QA command with stable UTF-8 diagnostics."""
    return subprocess.run(
        command,
        check=False,
        capture_output=True,
        env={**os.environ, "PYTHONUTF8": "1"},
        text=True,
        encoding="utf-8",
        errors="replace",
    )


def convert_pptx(pptx_path: Path, output_directory: Path) -> Path:
    """Convert a PowerPoint to PDF for deterministic slide inspection."""
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice is required to inspect PowerPoint output")
    result = run(
        [
            executable,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_directory),
            str(pptx_path),
        ]
    )
    converted = output_directory / f"{pptx_path.stem}.pdf"
    if result.returncode != 0 or not converted.is_file():
        raise RuntimeError(
            "PowerPoint conversion failed: "
            + (result.stderr or result.stdout).strip()
        )
    return converted


def contact_sheet(pdf_paths: list[tuple[str, Path]], output_path: Path) -> None:
    """Rasterize every page/slide into one compact review image."""
    if not shutil.which("pdftoppm"):
        raise RuntimeError("pdftoppm is required to create the QA contact sheet")
    try:
        from PIL import Image, ImageDraw
    except ImportError as error:
        raise RuntimeError("Pillow is required to create the QA contact sheet") from error
    images: list[tuple[str, object]] = []
    with tempfile.TemporaryDirectory() as directory:
        temporary = Path(directory)
        for label, pdf_path in pdf_paths:
            prefix = temporary / label
            result = run(
                [
                    "pdftoppm",
                    "-r",
                    "60",
                    "-png",
                    str(pdf_path),
                    str(prefix),
                ]
            )
            if result.returncode != 0:
                raise RuntimeError(
                    f"Unable to rasterize {pdf_path.name}: "
                    + (result.stderr or result.stdout).strip()
                )
            for index, path in enumerate(
                sorted(temporary.glob(f"{label}-*.png")), start=1
            ):
                images.append((f"{label.upper()} {index}", Image.open(path).convert("RGB")))
        if not images:
            raise RuntimeError("No pages were generated for the contact sheet")
        columns = min(3, len(images))
        label_height = 22
        gap = 12
        width = max(image.width for _, image in images)
        height = max(image.height for _, image in images)
        rows = (len(images) + columns - 1) // columns
        sheet = Image.new(
            "RGB",
            (
                columns * width + (columns + 1) * gap,
                rows * (height + label_height) + (rows + 1) * gap,
            ),
            "white",
        )
        draw = ImageDraw.Draw(sheet)
        for index, (label, image) in enumerate(images):
            column = index % columns
            row = index // columns
            x = gap + column * (width + gap)
            y = gap + row * (height + label_height + gap)
            draw.text((x, y), label, fill="black")
            sheet.paste(image, (x, y + label_height))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        sheet.save(output_path)
        for _, image in images:
            image.close()


def fix_for(issue: str) -> str:
    """Translate a layout diagnostic into a concrete repair lever."""
    normalized = issue.casefold()
    if "orphan tail" in normalized or "ends with heading" in normalized:
        return "keep the short tail with its preceding block or remove the forced break"
    if "forced break not applied" in normalized:
        return "restore the configured forced page break before the named section"
    if "full; minimum" in normalized:
        return "remove the following forced break or combine adjacent short sections"
    if "section " in normalized and "spans" in normalized:
        return "shorten that section, reduce body table rows, or move detail to an appendix"
    if "hard ceiling" in normalized:
        return "reduce narrative or body table rows in the longest reported sections"
    if "cover shows text" in normalized or "cover lists" in normalized:
        return "keep meta.cover_facts to identifying facts and move status or coverage into the body"
    if "raw status values" in normalized or "raw chip object" in normalized:
        return "pass statuses as chip objects so the renderer shows the plain-language label"
    if "has ordinal" in normalized:
        return "correct the ordinal suffix in the source text"
    if "wider than the page" in normalized:
        return "drop or merge table columns, split the table, or shorten long unbroken values"
    if "templated phrase" in normalized:
        return "rewrite the sentence around the specific figure, subject and consequence"
    return "repair the named source section, then re-compose and rerun QA"


def slide_lines(slide: object) -> list[str]:
    """Return the non-empty text lines of one slide in shape order."""
    lines: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            lines.extend(
                paragraph.text.strip()
                for paragraph in shape.text_frame.paragraphs
                if paragraph.text.strip()
            )
        if getattr(shape, "has_table", False):
            lines.extend(
                cell.text.strip()
                for row in shape.table.rows
                for cell in row.cells
                if cell.text.strip()
            )
    return lines


def deck_reader_diagnostics(
    slides: list[list[str]],
    report: dict[str, object],
    branding: dict[str, object],
) -> list[dict[str, object]]:
    """Apply the title-page and reader-text gates to rendered slide text."""
    metadata = report.get("meta") if isinstance(report.get("meta"), dict) else {}
    failures = cover_fact_count_errors(metadata)
    if metadata and slides:
        failures.extend(
            cover_page_failures(
                slides[0],
                cover_allowed_text(metadata, branding),
                location="slide 1",
            )
        )
    appendix_titles = {
        str(section.get("title", "")).strip().casefold()
        for section in report.get("sections", [])
        if isinstance(section, dict)
        and str(section.get("title", "")).strip().casefold().startswith("appendix")
    }
    first_appendix = next(
        (
            number
            for number, lines in enumerate(slides, start=1)
            if number > 1 and any(line.casefold() in appendix_titles for line in lines)
        ),
        None,
    )
    exempt = (
        set(range(first_appendix, len(slides) + 1)) if first_appendix else set()
    )
    failures.extend(
        reader_text_failures(slides, raw_status_exempt_pages=exempt, location="slide")
    )
    return [
        {"severity": "error", "code": "pptx-reader-text", "message": failure}
        for failure in failures
    ]


def inspect_pptx(
    path: Path,
    budget: dict[str, object],
    report: dict[str, object] | None = None,
    branding: dict[str, object] | None = None,
) -> dict[str, object]:
    """Check editable PowerPoint theme/font surfaces, slide budgets and reader text."""
    try:
        from pptx import Presentation
    except ImportError as error:
        raise RuntimeError("python-pptx is required to inspect PowerPoint output") from error
    presentation = Presentation(path)
    diagnostics: list[dict[str, object]] = []
    if report is not None:
        diagnostics.extend(
            deck_reader_diagnostics(
                [slide_lines(slide) for slide in presentation.slides],
                report,
                branding or {},
            )
        )
    slide_count = len(presentation.slides)
    pptx_budget = budget.get("pptx")
    pptx_budget = pptx_budget if isinstance(pptx_budget, dict) else {}
    target = pptx_budget.get("target_slides")
    if isinstance(target, dict):
        minimum, maximum = target.get("min"), target.get("max")
        if isinstance(minimum, int) and slide_count < minimum:
            diagnostics.append(
                {
                    "severity": "warning",
                    "code": "slide-soft-target",
                    "message": f"slide count {slide_count} is below soft target {minimum}",
                }
            )
        if isinstance(maximum, int) and slide_count > maximum:
            diagnostics.append(
                {
                    "severity": "warning",
                    "code": "slide-soft-target",
                    "message": f"slide count {slide_count} exceeds soft target {maximum}",
                }
            )
    ceiling = pptx_budget.get("hard_max_slides")
    if isinstance(ceiling, int) and slide_count > ceiling:
        diagnostics.append(
            {
                "severity": "error",
                "code": "slide-hard-ceiling",
                "message": f"slide count {slide_count} exceeds hard ceiling {ceiling}",
            }
        )
    font_issues = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for text_run in paragraph.runs:
                    if text_run.text.strip() and text_run.font.name != FONT_FAMILY:
                        font_issues += 1
    if font_issues:
        diagnostics.append(
            {
                "severity": "error",
                "code": "pptx-font",
                "message": f"{font_issues} PowerPoint text run(s) do not name {FONT_FAMILY}",
            }
        )
    return {
        "schema_version": 1,
        "theme_version": THEME_VERSION,
        "status": (
            "fail"
            if any(item["severity"] == "error" for item in diagnostics)
            else "pass"
        ),
        "slide_count": slide_count,
        "diagnostics": diagnostics,
    }


def preferred_break_candidates(
    report: dict[str, object],
    layout_result: dict[str, object],
    *,
    threshold: float = 0.55,
) -> set[str]:
    """Return preferred sections whose preceding body page is sparse."""
    body_pages = {
        int(item["page"]): float(item["fill"])
        for item in layout_result.get("body_pages", [])
        if isinstance(item, dict)
        and isinstance(item.get("page"), int)
        and isinstance(item.get("fill"), (int, float))
    }
    rendered_sections = {
        str(item.get("number", "")).lstrip("0") or "0": item
        for item in layout_result.get("sections", [])
        if isinstance(item, dict)
    }
    candidates: set[str] = set()
    for index, section in enumerate(report.get("sections", []), start=1):
        if not isinstance(section, dict) or section.get("page_break") != "prefer":
            continue
        number = str(section.get("num", index)).lstrip("0") or "0"
        rendered = rendered_sections.get(number)
        if not isinstance(rendered, dict):
            continue
        start_page = rendered.get("start_page")
        if (
            isinstance(start_page, int)
            and start_page > 1
            and (
                body_pages.get(start_page - 1, 1.0) < threshold
                or body_pages.get(start_page, 1.0) < threshold
            )
        ):
            candidates.add(str(section.get("title", "")).strip())
    return candidates


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("report", type=Path)
    parser.add_argument("--format", choices=("pdf", "pptx", "both"), default="pdf")
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--budget", type=Path)
    parser.add_argument("--brand", type=Path)
    parser.add_argument("--keep-workdir", action="store_true")
    args = parser.parse_args()

    here = Path(__file__).resolve().parent
    renderer = here / "render.py"
    checker = here / "check_layout.py"
    base = args.out.with_suffix("") if args.out.suffix.lower() in {".pdf", ".pptx"} else args.out
    design_diagnostics: list[dict[str, str]] = []
    try:
        report = json.loads(args.report.read_text(encoding="utf-8"))
        if not isinstance(report, dict):
            raise LayoutPolicyError(["report root must be an object"])
        branding = load_branding(args.brand.resolve() if args.brand else None)
        budget = load_layout_budget(args.budget) if args.budget else {}
        prepared = prepare_report_for_render(report)
        effective_budget = resolve_layout_budget(budget, report) if budget else {}
        post_render_layout_contract = (
            build_post_render_layout_contract(prepared, effective_budget)
            if effective_budget
            else {"schema_version": 1, "forced_breaks": [], "appendix_fill": []}
        )
        content_errors = (
            validate_report_budget(prepared, effective_budget)
            if budget
            else []
        )
        design_diagnostics = (
            lint_report_design(prepared, effective_budget)
            if effective_budget and not content_errors
            else []
        )
        if content_errors or any(
            item["severity"] == "error" for item in design_diagnostics
        ):
            issues = list(dict.fromkeys(content_errors))
            if issues:
                print(f"FAIL — {len(issues)} pre-render budget issue(s)")
                for issue in issues:
                    print(f"  - {issue}")
            if design_diagnostics:
                print(
                    f"FAIL — {len(design_diagnostics)} pre-render design issue(s)"
                )
                for item in design_diagnostics:
                    location = item["section_id"]
                    if item["section_title"]:
                        location += f" ({item['section_title']})"
                    print(
                        f"  - {item['severity'].upper()} [{item['code']}] "
                        f"{location}: {item['message']}"
                    )
            qa_path = base.parent / f"{base.name}-qa.json"
            qa_path.parent.mkdir(parents=True, exist_ok=True)
            qa_path.write_text(
                json.dumps(
                    {
                        "schema_version": 1,
                        "theme_version": THEME_VERSION,
                        "status": "fail",
                        "report": str(args.report),
                        "budget": str(args.budget) if args.budget else None,
                        "pre_render_budget_issues": issues,
                        "design_diagnostics": design_diagnostics,
                        "post_render_layout_contract": post_render_layout_contract,
                        "outputs": {},
                    },
                    ensure_ascii=False,
                    indent=2,
                )
                + "\n",
                encoding="utf-8",
            )
            print(f"QA diagnostics: {qa_path}")
            return 1
    except (
        OSError,
        json.JSONDecodeError,
        LayoutPolicyError,
        ValueError,
    ) as error:
        print(f"ERROR — unable to preflight report: {error}", file=sys.stderr)
        return 2

    workdir_path = Path(tempfile.mkdtemp(prefix="gradient-report-qa-"))
    prepared_path = workdir_path / "prepared-report.json"
    prepared_path.write_text(
        json.dumps(prepared, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    render_command = [
        sys.executable,
        str(renderer),
        str(prepared_path),
        "--format",
        args.format,
        "--out",
        str(base),
    ]
    if args.brand:
        render_command.extend(("--brand", str(args.brand)))
    render_result = run(
        render_command
    )
    if render_result.stdout:
        print(render_result.stdout, end="")
    if render_result.stderr:
        print(render_result.stderr, end="", file=sys.stderr)
    if render_result.returncode != 0:
        if not args.keep_workdir:
            shutil.rmtree(workdir_path, ignore_errors=True)
        return render_result.returncode

    qa_results: dict[str, object] = {
        "schema_version": 1,
        "theme_version": THEME_VERSION,
        "report": str(args.report),
        "budget": str(args.budget) if args.budget else None,
        "branding": {
            key: value
            for key, value in branding.items()
            if key != "client_logo"
        },
        "design_diagnostics": design_diagnostics,
        "post_render_layout_contract": post_render_layout_contract,
        "outputs": {},
    }
    failures: list[str] = []
    final_render_stderr = render_result.stderr or ""
    try:
        if args.format in {"pdf", "both"}:
            pdf_path = base.with_suffix(".pdf")
            result_path = workdir_path / "pdf-layout.json"
            command = [
                sys.executable,
                str(checker),
                str(pdf_path),
                "--report",
                str(prepared_path),
                "--json",
                str(result_path),
            ]
            if args.budget:
                command.extend(("--budget", str(args.budget)))
            if args.brand:
                command.extend(("--brand", str(args.brand.resolve())))
            checked = run(command)
            print("PDF layout pass 1:")
            print((checked.stdout or checked.stderr).rstrip())
            if result_path.is_file():
                pdf_result = json.loads(result_path.read_text(encoding="utf-8"))
            elif checked.returncode != 0:
                return checked.returncode
            else:
                pdf_result = {}
            current = prepared
            all_removed: list[str] = []
            for attempt in range(1, MAX_REPAIR_PASSES + 1):
                if checked.returncode == 0:
                    break
                removable = preferred_break_candidates(current, pdf_result)
                if not removable:
                    break
                current, removed = repair_preferred_breaks(current, removable)
                if not removed:
                    break
                all_removed.extend(removed)
                repaired_path = workdir_path / "repaired-report.json"
                repaired_path.write_text(
                    json.dumps(current, ensure_ascii=False, indent=2) + "\n",
                    encoding="utf-8",
                )
                rerender_command = list(render_command)
                rerender_command[2] = str(repaired_path)
                rerendered = run(rerender_command)
                if rerendered.stdout:
                    print(rerendered.stdout, end="")
                if rerendered.stderr:
                    print(rerendered.stderr, end="", file=sys.stderr)
                if rerendered.returncode != 0:
                    return rerendered.returncode
                final_render_stderr = rerendered.stderr or ""
                command[4] = str(repaired_path)
                checked = run(command)
                print(f"Preferred-break repair {attempt}: " + ", ".join(removed))
                print("PDF layout after repair:")
                print((checked.stdout or checked.stderr).rstrip())
                if not result_path.is_file():
                    return checked.returncode or 2
                pdf_result = json.loads(
                    result_path.read_text(encoding="utf-8")
                )
            qa_results["repair"] = {
                "rerendered": bool(all_removed),
                "removed_preferred_breaks": all_removed,
            }
            qa_results["outputs"]["pdf"] = pdf_result
            failures.extend(
                str(item["message"])
                for item in pdf_result.get("diagnostics", [])
                if item.get("severity") == "error"
            )
            failures.extend(
                line.removeprefix("WARNING: ").strip()
                for line in final_render_stderr.splitlines()
                if "wider than the page" in line
            )
            pdf_contact_path = (
                base.parent / f"{base.name}-pdf-contact-sheet.png"
            )
            contact_sheet([("pdf", pdf_path)], pdf_contact_path)
            print(f"PDF contact sheet: {pdf_contact_path}")
        if args.format in {"pptx", "both"}:
            pptx_path = base.with_suffix(".pptx")
            pptx_result = inspect_pptx(
                pptx_path,
                effective_budget,
                prepared,
                branding,
            )
            qa_results["outputs"]["pptx"] = pptx_result
            failures.extend(
                str(item["message"])
                for item in pptx_result["diagnostics"]
                if item["severity"] == "error"
            )
            converted = convert_pptx(
                pptx_path, workdir_path
            )
            pptx_contact_path = (
                base.parent / f"{base.name}-pptx-contact-sheet.png"
            )
            contact_sheet([("pptx", converted)], pptx_contact_path)
            print(f"PowerPoint contact sheet: {pptx_contact_path}")
        qa_results["status"] = "fail" if failures else "pass"
        qa_path = base.parent / f"{base.name}-qa.json"
        qa_path.write_text(
            json.dumps(qa_results, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"QA diagnostics: {qa_path}")
        if failures:
            print("Prioritized repairs:")
            for index, issue in enumerate(failures, start=1):
                print(f"  {index}. {issue} — {fix_for(issue)}")
            return 1
        print("PASS — requested outputs passed layout QA")
        return 0
    except RuntimeError as error:
        print(f"ERROR — {error}", file=sys.stderr)
        return 2
    finally:
        if args.keep_workdir:
            print(f"QA work directory: {workdir_path}")
        else:
            shutil.rmtree(workdir_path, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
