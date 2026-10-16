#!/usr/bin/env python3
"""Editable 16:9 PowerPoint renderer for validated Gradient report JSON.

The public entry point is ``render_report(report, output_path)``.  The module
imports python-pptx lazily so validation and CLI diagnostics remain useful when
the optional PowerPoint dependency is not installed.
"""

from __future__ import annotations

import json
import math
import re
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any, TypeAlias, cast
from xml.sax.saxutils import escape as xml_escape

from report_layout import (
    LayoutPolicyError,
    effective_chart_mode,
    prepare_report_for_render,
)
from report_theme import (
    AMBER as THEME_AMBER,
    BRAND_NAME,
    CORAL as THEME_CORAL,
    COVER_FIELD_LIMITS,
    DEFAULT_COVER_DISCLOSURE,
    FONT_FAMILY,
    INK as THEME_INK,
    INK2 as THEME_INK_2,
    LIME as THEME_LIME,
    LIME_DK as THEME_LIME_DARK,
    MUTED as THEME_MUTED,
    PANEL as THEME_PANEL,
    PAPER as THEME_PAPER,
    PIE as THEME_PIE,
    PPTX_SLIDE_HEIGHT_IN,
    PPTX_SLIDE_WIDTH_IN,
    RULE as THEME_RULE,
    SERIES as THEME_SERIES,
    SIGNAL as THEME_SIGNAL,
    SLATE as THEME_SLATE,
    STATUS_COLOR as THEME_STATUS_COLOR,
    STATUS_LABEL as THEME_STATUS_LABEL,
    THEME_VERSION,
    WASH as THEME_WASH,
    default_confidentiality,
    pptx_hex,
    status_key,
    status_label,
    validate_potx_theme,
    without_lead,
)


JsonValue: TypeAlias = (
    str | int | float | bool | None | list["JsonValue"] | dict[str, "JsonValue"]
)
JsonObject: TypeAlias = dict[str, JsonValue]

INK = pptx_hex(THEME_INK)
PANEL = pptx_hex(THEME_PANEL)
LIME = pptx_hex(THEME_LIME)
LIME_DARK = pptx_hex(THEME_LIME_DARK)
INK_2 = pptx_hex(THEME_INK_2)
MUTED = pptx_hex(THEME_MUTED)
RULE = pptx_hex(THEME_RULE)
PAPER = pptx_hex(THEME_PAPER)
WASH = pptx_hex(THEME_WASH)
AMBER = pptx_hex(THEME_AMBER)
CORAL = pptx_hex(THEME_CORAL)
SLATE = pptx_hex(THEME_SLATE)
SERIES = tuple(pptx_hex(color) for color in THEME_SERIES)
PIE = tuple(pptx_hex(color) for color in THEME_PIE)
SLIDE_WIDTH_IN = PPTX_SLIDE_WIDTH_IN
SLIDE_HEIGHT_IN = PPTX_SLIDE_HEIGHT_IN
TABLE_ROWS_PER_SLIDE = 12
SLIDE_BODY_CAPACITY = 5.15
SLIDE_BODY_WIDTH = 11.93
SLIDE_BODY_HEIGHT = 5.35
SLIDE_COLUMN_WIDTH = 5.73
STACK_GAP = 0.12
VISUAL_MIN_HEIGHT = 3.0
ALL_SOURCE_TAGS = re.compile(r"(?<![A-Z0-9])S(\d+)")
PLACEHOLDERS = (
    re.compile(r"<[^<>\n]{2,100}>"),
    re.compile(r"\{\{[^{}\n]+\}\}"),
    re.compile(r"\b(?:TBD|TODO|FIXME|TK)\b", re.IGNORECASE),
    re.compile(r"\[(?:INSERT|PLACEHOLDER)[^\]]*\]", re.IGNORECASE),
)
LAYOUTS = frozenset(
    {"cover", "divider", "exhibit", "two-up", "table", "text", "appendix"}
)
VISUAL_BLOCKS = frozenset(
    {"bars", "pie", "percentiles", "line", "stacked", "waterfall", "band", "heat", "timeline", "chart"}
)
TABLE_BLOCKS = frozenset({"table", "kv", "coverage"})
EXHIBIT_COMPANIONS = frozenset({"text", "markdown", "callout", "statement"})
# Heat maps, timelines and table-mode charts need the full slide width.
PAIRABLE_VISUALS = frozenset({"bars", "pie", "percentiles", "line", "stacked", "waterfall", "band"})
SUPPORTED_BLOCKS = frozenset(
    {
        "text",
        "markdown",
        "bullets",
        "kv",
        "table",
        "tiles",
        "bars",
        "pie",
        "percentiles",
        "line",
        "stacked",
        "waterfall",
        "band",
        "heat",
        "timeline",
        "callout",
        "coverage",
        "findings",
        "questions",
        "statement",
        "chart",
        "two_col",
        "pagebreak",
    }
)


class RenderValidationError(RuntimeError):
    """Raised when a deck would contain invalid or overflowing content."""

    def __init__(self, errors: list[str]) -> None:
        super().__init__("\n".join(errors))
        self.errors = errors


def _object(value: JsonValue | object) -> JsonObject:
    return cast(JsonObject, value) if isinstance(value, dict) else {}


def _array(value: JsonValue | object) -> list[JsonValue]:
    return cast(list[JsonValue], value) if isinstance(value, list) else []


def _text(value: JsonValue | object, default: str = "") -> str:
    return str(value) if value is not None else default


def _number(value: JsonValue | object, default: float = 0.0) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        amount = float(value)
        return amount if math.isfinite(amount) else default
    return default


def _axis_number_format(
    maximum: float,
    *,
    decimals: int = 1,
    currency: str = "",
    value_format: str = "",
    suffix: str = "",
) -> str:
    """Return a compact PowerPoint axis format matching the PDF renderer."""
    prefix = {
        "USD": "$",
        "EUR": "€",
        "GBP": "£",
        "JPY": "¥",
    }.get(currency.upper(), f"{currency} " if currency else "")
    if value_format != "currency":
        prefix = ""
    scale = ""
    marker = ""
    absolute = abs(maximum)
    if absolute >= 1_000_000_000:
        scale, marker = ",,,", "B"
    elif absolute >= 1_000_000:
        scale, marker = ",,", "M"
    elif absolute >= 1_000:
        scale, marker = ",", "K"
    precision = 0 if marker and absolute / (1_000 ** len(scale)) >= 100 else decimals
    zeroes = "0" if precision <= 0 else "0." + ("0" * precision)
    literal_suffix = f'"{marker}{suffix}"' if marker or suffix else ""
    return f'{prefix}{zeroes}{scale}{literal_suffix}'


def _plain(value: JsonValue | object) -> str:
    text = _text(value)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\w)\*([^*]+?)\*(?!\w)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.strip()


def _ordinal(value: float) -> str:
    number = int(round(value))
    suffix = "th"
    if not 10 <= abs(number) % 100 <= 20:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(abs(number) % 10, "th")
    return f"{number}{suffix}"


def _compact(value: JsonValue | object, words: int = 28) -> str:
    tokens = _plain(value).split()
    if len(tokens) <= words:
        return " ".join(tokens)
    return " ".join(tokens[:words]).rstrip(".,;:") + "…"

def _cover_disclosure(meta: JsonObject) -> str:
    explicit = _text(meta.get("cover_disclosure")).strip()
    if explicit:
        return explicit
    return DEFAULT_COVER_DISCLOSURE if _text(meta.get("data_scope_banner")).strip() else ""


def _source_tags(value: JsonValue | object) -> list[str]:
    encoded = json.dumps(value, ensure_ascii=False)
    found = {int(number) for number in ALL_SOURCE_TAGS.findall(encoded)}
    return [f"S{number}" for number in sorted(found)]


def _tone_color(tone: str) -> str:
    return {
        "good": LIME,
        "watch": AMBER,
        "warning": AMBER,
        "bad": CORAL,
        "info": INK,
    }.get(tone.lower(), SLATE)


def _status_color(status: str) -> str:
    raw = status.lower()
    color = THEME_STATUS_COLOR.get(raw) or THEME_STATUS_COLOR.get(
        raw.replace("_", " ")
    )
    return pptx_hex(color) if color else SLATE

def _status_display(status: str, label: str = "", display_label: str = "") -> str:
    """Return the reader-facing label while retaining status for color semantics."""
    raw = status.strip() or label.strip()
    authored = label.strip()
    if display_label.strip() or not authored or status_key(authored) == status_key(raw):
        return status_label(raw, display_label)
    known = THEME_STATUS_LABEL.get(status_key(authored))
    if known:
        return known
    return authored[:1].upper() + authored[1:]


LEGACY_CHIP_TEXT = re.compile(
    r"""^\s*\{\s*['"]chip['"]\s*:\s*['"]([^'"]*)['"]\s*"""
    r"""(?:,\s*['"]status['"]\s*:\s*['"]([^'"]*)['"]\s*)?"""
    r"""(?:,\s*['"]display_label['"]\s*:\s*['"]([^'"]*)['"]\s*)?\}\s*$"""
)


def _chip_parts(value: JsonValue | object) -> tuple[str, str, str] | None:
    """Return (label, status, display_label) for a chip object or its repr text."""
    object_value = _object(value)
    if object_value.get("chip") is not None:
        label = _text(object_value.get("chip"))
        return (
            label,
            _text(object_value.get("status") or label),
            _text(object_value.get("display_label")),
        )
    if isinstance(value, str):
        match = LEGACY_CHIP_TEXT.fullmatch(value)
        if match:
            return match.group(1), match.group(2) or match.group(1), match.group(3) or ""
    return None

def _bars_mark(block: JsonObject) -> str:
    mark = _text(block.get("mark") or block.get("style") or "bars").strip().casefold()
    return "dots" if mark in {"dots", "paired-dot", "paired_dot"} else "bars"

def _normalized_bars(
    block: JsonObject,
) -> tuple[list[JsonObject], list[JsonObject], bool]:
    """Normalize validated legacy or named-series bars for rendering."""
    raw_items = [_object(value) for value in _array(block.get("items"))]
    raw_series = block.get("series")
    if raw_series is None:
        items: list[JsonObject] = []
        for item in raw_items:
            value = _number(item.get("value"))
            items.append({
                **item,
                "values": [value],
                "displays": [item.get("display", item.get("value", value))],
                "_authored_displays": item.get("display") is not None,
            })
        return [{"name": "Value"}], items, False
    series = [_object(value) for value in _array(raw_series)]
    items = []
    for item in raw_items:
        values = [_number(value) for value in _array(item.get("values"))]
        displays = (
            _array(item.get("displays"))
            if item.get("displays") is not None
            else cast(list[JsonValue], values)
        )
        items.append({
            **item,
            "values": values,
            "displays": displays,
            "_authored_displays": item.get("displays") is not None,
        })
    return series, items, True

def _bars_validation_errors(block: JsonObject, prefix: str) -> list[str]:
    errors: list[str] = []
    items = [_object(value) for value in _array(block.get("items"))]
    raw_series = block.get("series")
    if raw_series is None:
        if any(
            not isinstance(item.get("value"), (int, float))
            or isinstance(item.get("value"), bool)
            or not math.isfinite(float(cast(float, item.get("value"))))
            for item in items
        ):
            errors.append(f"{prefix} requires finite numeric values")
        if _bars_mark(block) == "dots":
            errors.append(f"{prefix} paired-dot mode requires named series")
        return errors
    series = [_object(value) for value in _array(raw_series)]
    if not series or any(not _text(item.get("name")).strip() for item in series):
        errors.append(f"{prefix} requires non-empty named series")
        return errors
    for item_index, item in enumerate(items, start=1):
        values = _array(item.get("values"))
        if len(values) != len(series):
            errors.append(
                f"{prefix} item {item_index} values length must match series"
            )
            continue
        if any(
            not isinstance(value, (int, float))
            or isinstance(value, bool)
            or not math.isfinite(float(value))
            for value in values
        ):
            errors.append(f"{prefix} item {item_index} requires finite values")
        displays = item.get("displays")
        if displays is not None and len(_array(displays)) != len(series):
            errors.append(
                f"{prefix} item {item_index} displays length must match series"
            )
    if _bars_mark(block) == "dots" and len(series) != 2:
        errors.append(f"{prefix} paired-dot mode requires exactly two series")
    return errors

def _bar_gap_display(item: JsonObject, difference: float, block: JsonObject) -> str:
    authored = item.get("gap_display")
    if authored is not None:
        return _text(authored)
    decimals = max(0, int(_number(block.get("decimals"), 1)))
    suffix = _text(block.get("suffix") or block.get("y_suffix"))
    label = _text(block.get("gap_label"), "Gap")
    rounded = round(difference, decimals)
    if rounded == 0:
        rounded = 0.0
    return f"{label} {rounded:+.{decimals}f}{suffix}"

def _set_native_data_label_text(chart_series: Any, labels: list[str]) -> None:
    """Attach editable custom text labels to one native chart series."""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn

    series_element = chart_series._element
    existing = series_element.find(qn("c:dLbls"))
    if existing is not None:
        series_element.remove(existing)
    label_xml = "".join(
        (
            f'<c:dLbl><c:idx val="{index}"/><c:tx><c:rich>'
            '<a:bodyPr/><a:lstStyle/><a:p><a:r>'
            f'<a:rPr lang="en-US" sz="800"/><a:t>{xml_escape(label)}</a:t>'
            '</a:r></a:p></c:rich></c:tx><c:dLblPos val="outEnd"/>'
            '<c:showLegendKey val="0"/><c:showVal val="0"/>'
            '<c:showCatName val="0"/><c:showSerName val="0"/>'
            '<c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>'
        )
        for index, label in enumerate(labels)
    )
    data_labels = parse_xml(
        f'<c:dLbls {nsdecls("c", "a")}>{label_xml}'
        '<c:dLblPos val="outEnd"/><c:showLegendKey val="0"/>'
        '<c:showVal val="0"/><c:showCatName val="0"/>'
        '<c:showSerName val="0"/><c:showPercent val="0"/>'
        '<c:showBubbleSize val="0"/></c:dLbls>'
    )
    insert_at = len(series_element)
    for index, child in enumerate(series_element):
        if child.tag in {qn("c:cat"), qn("c:val")}:
            insert_at = index
            break
    series_element.insert(insert_at, data_labels)


def _collect_strings(value: JsonValue) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, list):
        return [item for child in value for item in _collect_strings(child)]
    if isinstance(value, dict):
        return [item for child in value.values() for item in _collect_strings(child)]
    return []


def validate_report_shape(report: JsonObject) -> list[str]:
    """Validate render-critical structure without replacing skill validation."""
    errors: list[str] = []
    meta = _object(report.get("meta"))
    if not _text(meta.get("title")).strip():
        errors.append("meta.title is required")
    sections = _array(report.get("sections"))
    slides = _array(report.get("slides"))
    if not sections and not slides:
        errors.append("report must contain a non-empty sections or slides array")
    for text in _collect_strings(report):
        for pattern in PLACEHOLDERS:
            match = pattern.search(text)
            if match is not None:
                errors.append(f"unresolved placeholder: {match.group(0)}")
    containers = sections if sections else slides
    for index, value in enumerate(containers, start=1):
        container = _object(value)
        if not _text(container.get("title")).strip():
            errors.append(f"content item {index} has an empty title")
        layout = _text(container.get("layout"))
        aliases = {"full", "wide", "two", "section"}
        if layout and layout not in LAYOUTS and layout not in aliases:
            errors.append(f"content item {index} has unsupported layout '{layout}'")
        for block_index, block_value in enumerate(_blocks_for_container(container), start=1):
            block = _object(block_value)
            block_type = _text(block.get("type"))
            prefix = f"content item {index}, block {block_index} ({block_type or 'unknown'})"
            if not block_type:
                errors.append(f"{prefix} has no type")
                continue
            if block_type not in SUPPORTED_BLOCKS:
                errors.append(f"{prefix} has unsupported block type '{block_type}'")
                continue
            errors.extend(_validate_block_data(block, prefix))
    return list(dict.fromkeys(errors))


def _validate_block_data(block: JsonObject, prefix: str) -> list[str]:
    errors: list[str] = []
    block_type = _text(block.get("type"))
    required_arrays = {
        "bullets": "items",
        "kv": "rows",
        "table": "rows",
        "tiles": "tiles",
        "bars": "items",
        "pie": "items",
        "percentiles": "items",
        "coverage": "items",
        "questions": "items",
        "stacked": "items",
        "waterfall": "items",
        "band": "items",
        "heat": "rows",
        "timeline": "items",
    }
    key = required_arrays.get(block_type)
    if key is not None and not _array(block.get(key)):
        if block_type == "findings" and block.get("empty_title"):
            return errors
        errors.append(f"{prefix} has no data in {key}")
    if block_type == "line":
        series = [_object(value) for value in _array(block.get("series"))]
        if not series or not any(_array(item.get("points")) for item in series):
            errors.append(f"{prefix} has no line points")
    if block_type == "bars":
        errors.extend(_bars_validation_errors(block, prefix))
    if block_type == "pie":
        values = [
            value
            for item in (_object(entry) for entry in _array(block.get("items")))
            if isinstance(value := item.get("value"), (int, float))
            and not isinstance(value, bool)
            and math.isfinite(value)
        ]
        if len(values) != len(_array(block.get("items"))):
            errors.append(f"{prefix} requires finite numeric values")
        elif any(value < 0 for value in values):
            errors.append(f"{prefix} requires nonnegative values")
        elif not any(value > 0 for value in values):
            errors.append(f"{prefix} requires at least one positive slice")
        elif sum(value > 0 for value in values) > len(PIE):
            errors.append(f"{prefix} supports at most {len(PIE)} positive slices")
    if block_type == "chart":
        chart = _object(block.get("chart"))
        status = _text(chart.get("status"))
        if status == "ok" and not _array(chart.get("rows")):
            errors.append(f"{prefix} has status ok but no rows")
        if status != "ok" and not _text(chart.get("error") or status):
            errors.append(f"{prefix} has unavailable data without a reason")
    if block_type in {"text", "statement", "callout"}:
        if not _text(block.get("text")).strip():
            errors.append(f"{prefix} has empty text")
    if block_type == "findings":
        if not _array(block.get("items")) and not _text(block.get("empty_title")).strip():
            errors.append(f"{prefix} has no findings and no empty-state label")
    if block_type == "two_col":
        if not _array(block.get("left")) and not _array(block.get("right")):
            errors.append(f"{prefix} has two empty columns")
    return errors


def _blocks_for_container(container: JsonObject) -> list[JsonValue]:
    blocks = _array(container.get("blocks"))
    if blocks:
        return blocks
    left = _array(container.get("left"))
    right = _array(container.get("right"))
    if left or right:
        return [{"type": "two_col", "left": left, "right": right}]
    return []


def _split_table_block(block: JsonObject) -> list[JsonObject]:
    block_type = _text(block.get("type"))
    if block_type == "chart":
        chart = _object(block.get("chart"))
        if (
            _text(chart.get("status")) == "ok"
            and effective_chart_mode(chart) == "table"
            and len(_array(chart.get("rows"))) > TABLE_ROWS_PER_SLIDE
        ):
            rows = _array(chart.get("rows"))
            chunks: list[JsonObject] = []
            title = _text(chart.get("title"), "Chart data")
            for start in range(0, len(rows), TABLE_ROWS_PER_SLIDE):
                clone = deepcopy(block)
                clone_chart = _object(clone.get("chart"))
                clone_chart["rows"] = rows[start : start + TABLE_ROWS_PER_SLIDE]
                if start:
                    clone_chart["title"] = f"{title} — continued"
                chunks.append(clone)
            return chunks
    if block_type != "table":
        return [block]
    rows = _array(block.get("rows"))
    if len(rows) <= TABLE_ROWS_PER_SLIDE:
        return [block]
    chunks: list[JsonObject] = []
    title = _text(block.get("title"), "Table")
    for start in range(0, len(rows), TABLE_ROWS_PER_SLIDE):
        clone = deepcopy(block)
        clone["rows"] = rows[start : start + TABLE_ROWS_PER_SLIDE]
        if start:
            clone["title"] = f"{title} — continued"
        clone["_continuation"] = start // TABLE_ROWS_PER_SLIDE + 1
        chunks.append(clone)
    return chunks


def _markdown_parts(raw: str) -> tuple[list[str], list[str]]:
    """Split slide markdown into bullet items and plain paragraphs."""
    items: list[str] = []
    paragraphs: list[str] = []
    for line in (line.strip() for line in raw.splitlines()):
        if not line:
            continue
        normalized = re.sub(r"^#{1,6}\s+", "", line)
        normalized = re.sub(r"^\d+[.)]\s+", "", normalized)
        if line.startswith(("-", "*", "•")) or re.match(r"^\d+[.)]\s+", line):
            items.append(normalized.lstrip("-*• "))
        elif not re.match(r"^\|[-:| ]+\|?$", line):
            paragraphs.append(normalized.replace("|", " · "))
    return items, paragraphs


def _wrapped_height(
    paragraphs: list[str], size: float, width: float, paragraph_gap_pt: float = 0.0
) -> float:
    """Estimate text height with the same wrap model as the overflow guard plus a safety margin."""
    characters_per_line = max(1, int(width * 72 / (size * 0.52)))
    lines = sum(
        max(1, math.ceil(len(_plain(paragraph)) / characters_per_line))
        for paragraph in paragraphs
    )
    gaps = paragraph_gap_pt * max(0, len(paragraphs) - 1)
    return 1.25 * (lines * size * 1.2 + gaps) / 72


def _estimate_block(block: JsonObject, width: float = SLIDE_BODY_WIDTH) -> float:
    block_type = _text(block.get("type"))
    title_height = 0.3 if block.get("title") else 0.0
    if block_type == "text":
        return max(0.5, title_height + _wrapped_height([_text(block.get("text"))], 11, width))
    if block_type == "markdown":
        items, paragraphs = _markdown_parts(_text(block.get("text")))
        if items:
            return max(0.6, _wrapped_height(items, 10.5, width, 5))
        # Paragraphs render joined by a blank line.
        return max(0.6, _wrapped_height(paragraphs, 11, width, 13.2))
    if block_type == "bullets":
        items = [_plain(value) for value in _array(block.get("items"))]
        return max(0.5, title_height + _wrapped_height(items, 10.5, width, 5))
    if block_type in {"table", "kv"}:
        return 0.55 + 0.34 * len(_array(block.get("rows")))
    if block_type == "tiles":
        return 1.25 if len(_array(block.get("tiles"))) <= 4 else 2.5
    if block_type in VISUAL_BLOCKS:
        return 3.9
    if block_type in {"coverage", "findings", "questions"}:
        return 0.55 + 0.55 * len(_array(block.get("items")))
    if block_type == "callout":
        text = _plain(block.get("text"))
        if block.get("role") == "analysis":
            text = _compact(text)
        offset = 0.43 if block.get("title") else 0.15
        return max(0.75, offset + 0.1 + _wrapped_height([text], 9.5, width - 0.32))
    if block_type == "statement":
        return max(0.8, len(_plain(block.get("text"))) / 260.0)
    if block_type == "two_col":
        column = (width - 0.47) / 2
        return max(
            sum(_estimate_block(_object(value), column) for value in _array(block.get("left"))),
            sum(_estimate_block(_object(value), column) for value in _array(block.get("right"))),
        )
    return 1.0


def _plan_section(section: JsonObject) -> list[tuple[str, list[JsonObject]]]:
    requested = _text(section.get("layout"))
    if requested == "section":
        requested = "divider"
    elif requested in {"two", "wide"}:
        requested = "two-up"
    elif requested == "full":
        requested = ""
    blocks: list[JsonObject] = []
    for value in _blocks_for_container(section):
        blocks.extend(_split_table_block(_object(value)))
    if requested == "divider":
        plans: list[tuple[str, list[JsonObject]]] = [("divider", [])]
        if not blocks:
            return plans
        requested = ""
    else:
        plans = []
    if not blocks:
        return plans or [("text", [])]
    current: list[JsonObject] = []
    current_height = 0.0

    def flush() -> None:
        nonlocal current, current_height
        if current:
            plans.append((requested or _infer_layout(section, current), current))
        current = []
        current_height = 0.0

    index = 0
    while index < len(blocks):
        block = blocks[index]
        block_type = _text(block.get("type"))
        weight = _estimate_block(block)
        index += 1
        if block_type == "pagebreak":
            flush()
            continue
        if block_type in VISUAL_BLOCKS:
            # An exhibit shares its slide only with a short lead-in and the commentary that reads it.
            exhibit = [block]
            if (
                current
                and _text(current[-1].get("type")) in EXHIBIT_COMPANIONS
                and _exhibit_fits(exhibit + current[-1:])
            ):
                lead = current.pop()
                current_height -= _estimate_block(lead)
                exhibit.insert(0, lead)
            flush()
            while index < len(blocks):
                follow = blocks[index]
                follow_type = _text(follow.get("type"))
                leads_next_visual = (
                    follow_type in {"text", "markdown"}
                    and index + 1 < len(blocks)
                    and _text(blocks[index + 1].get("type")) in VISUAL_BLOCKS
                )
                if (
                    follow_type not in EXHIBIT_COMPANIONS
                    or leads_next_visual
                    or not _exhibit_fits(exhibit + [follow])
                ):
                    break
                exhibit.append(follow)
                index += 1
            plans.append((requested or _infer_layout(section, [block]), exhibit))
            continue
        isolate = block_type == "two_col" or (
            block_type == "table" and len(_array(block.get("rows"))) > 6
        )
        if isolate or current_height + weight > SLIDE_BODY_CAPACITY:
            flush()
        current.append(block)
        current_height += weight
        if isolate:
            plans.append((requested or _infer_layout(section, [block]), current))
            current = []
            current_height = 0.0
    flush()
    return plans if requested else _pair_exhibits(plans)


def _exhibit_fits(blocks: list[JsonObject], width: float = SLIDE_BODY_WIDTH) -> bool:
    """Return whether the companions leave each visual its minimum height."""
    visuals = sum(1 for block in blocks if _text(block.get("type")) in VISUAL_BLOCKS)
    companions = sum(
        _estimate_block(block, width)
        for block in blocks
        if _text(block.get("type")) not in VISUAL_BLOCKS
    )
    gaps = STACK_GAP * (len(blocks) - 1)
    return companions + gaps + VISUAL_MIN_HEIGHT * visuals <= SLIDE_BODY_HEIGHT


def _pairable_exhibit(plan: tuple[str, list[JsonObject]]) -> bool:
    layout, blocks = plan
    visuals = [block for block in blocks if _text(block.get("type")) in VISUAL_BLOCKS]
    return (
        layout == "exhibit"
        and len(visuals) == 1
        and _text(visuals[0].get("type")) in PAIRABLE_VISUALS
        and _exhibit_fits(blocks, SLIDE_COLUMN_WIDTH)
    )


def _pair_exhibits(
    plans: list[tuple[str, list[JsonObject]]]
) -> list[tuple[str, list[JsonObject]]]:
    """Set consecutive simple exhibits from one section side by side."""
    paired: list[tuple[str, list[JsonObject]]] = []
    index = 0
    while index < len(plans):
        if (
            index + 1 < len(plans)
            and _pairable_exhibit(plans[index])
            and _pairable_exhibit(plans[index + 1])
        ):
            left, right = plans[index][1], plans[index + 1][1]
            paired.append(("two-up", [{"type": "two_col", "left": left, "right": right}]))
            index += 2
            continue
        paired.append(plans[index])
        index += 1
    return paired


def _infer_layout(section: JsonObject, blocks: list[JsonObject]) -> str:
    title = _text(section.get("title")).lower()
    if title.startswith("appendix"):
        return "appendix"
    types = {_text(block.get("type")) for block in blocks}
    if "two_col" in types:
        return "two-up"
    if types & VISUAL_BLOCKS:
        return "exhibit"
    if types and types.issubset(TABLE_BLOCKS):
        return "table"
    return "text"


class PptxRenderer:
    """Render one validated report object to an editable PowerPoint deck."""

    def __init__(
        self,
        report: JsonObject,
        template_path: Path | None = None,
        branding: JsonObject | None = None,
    ) -> None:
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError as error:
            raise RuntimeError(
                "PowerPoint rendering requires python-pptx (pip install python-pptx)"
            ) from error
        self._pptx: Any = __import__("pptx")
        if template_path and template_path.suffix.casefold() == ".potx":
            theme_errors = validate_potx_theme(template_path)
            if theme_errors:
                raise RenderValidationError(theme_errors)
            source = BytesIO(template_path.read_bytes())
            converted = BytesIO()
            with zipfile.ZipFile(source, "r") as archive, zipfile.ZipFile(
                converted, "w"
            ) as output:
                for info in archive.infolist():
                    payload = archive.read(info.filename)
                    if info.filename == "[Content_Types].xml":
                        payload = payload.replace(
                            b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                            b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                        )
                    output.writestr(info, payload)
            converted.seek(0)
            self._presentation = Presentation(converted)
        else:
            self._presentation = (
                Presentation(str(template_path)) if template_path else Presentation()
            )
        self._presentation.slide_width = Inches(SLIDE_WIDTH_IN)
        self._presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
        self._blank_layout: Any = self._presentation.slide_layouts[6]
        try:
            self._report = cast(JsonObject, prepare_report_for_render(report))
        except LayoutPolicyError as error:
            raise RenderValidationError(error.errors) from error
        self._meta = _object(self._report.get("meta"))
        self._branding = branding or {}
        self._notes: dict[int, list[str]] = {}
        self._layout_errors: list[str] = []
        self._slide_number = 0

    def render(self, output_path: Path) -> None:
        """Build and save the presentation, raising on render validation."""
        errors = validate_report_shape(self._report)
        if errors:
            raise RenderValidationError(errors)
        self._add_cover()
        source_containers = _array(self._report.get("sections"))
        if source_containers:
            self._render_sections([_object(value) for value in source_containers])
        else:
            self._render_deck_slides(
                [_object(value) for value in _array(self._report.get("slides"))]
            )
        self._write_notes()
        self._apply_run_fonts()
        errors = self._validate_presentation()
        if errors:
            raise RenderValidationError(errors)
        self._presentation.core_properties.title = _text(self._meta.get("title"))
        self._presentation.core_properties.author = (
            _text(self._branding.get("client_name")) or "GradientCIO"
        )
        self._presentation.core_properties.subject = _text(
            self._meta.get("header_label"), "Gradient report"
        )
        self._presentation.core_properties.keywords = (
            f"Gradient report theme {THEME_VERSION}"
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self._presentation.save(str(output_path))

    def _apply_run_fonts(self) -> None:
        """Name the licensed report font on every generated text run."""
        def apply_frame(frame: Any) -> None:
            for paragraph in frame.paragraphs:
                paragraph.font.name = FONT_FAMILY
                for run in paragraph.runs:
                    run.font.name = FONT_FAMILY

        for slide in self._presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    apply_frame(shape.text_frame)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            apply_frame(cell.text_frame)

    def _new_slide(self) -> Any:
        slide = self._presentation.slides.add_slide(self._blank_layout)
        self._slide_number += 1
        return slide

    def _add_cover(self) -> None:
        slide = self._new_slide()
        self._shape(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, INK, INK)
        self._shape(slide, 10.25, 0.18, 2.75, 2.75, PANEL, PANEL, "oval")
        client_name = _text(self._branding.get("client_name"))
        logo = _text(self._branding.get("client_logo"))
        if logo and Path(logo).suffix.casefold() in {".png", ".jpg", ".jpeg"}:
            from pptx.util import Inches

            slide.shapes.add_picture(
                logo,
                Inches(0.75),
                Inches(0.45),
                height=Inches(0.42),
            )
        else:
            self._text_box(
                slide,
                0.75,
                0.5,
                5.2,
                0.35,
                client_name or "GradientCIO",
                12,
                PAPER,
                bold=True,
            )
        if client_name and bool(self._branding.get("powered_by", True)):
            self._text_box(
                slide,
                9.5,
                0.54,
                3.1,
                0.25,
                "Powered by GradientCIO",
                7,
                MUTED,
                align="right",
            )
        self._text_box(
            slide,
            0.75,
            1.82,
            10.9,
            0.3,
            _text(self._meta.get("eyebrow") or self._meta.get("header_label")).upper(),
            8.5,
            LIME,
            bold=True,
        )
        self._text_box(
            slide,
            0.75,
            2.2,
            10.9,
            1.35,
            _text(self._meta.get("title")),
            33,
            PAPER,
            bold=True,
        )
        subtitle = _text(self._meta.get("subtitle"))
        if subtitle:
            self._text_box(slide, 0.75, 3.58, 10.6, 0.65, subtitle, 14, SLATE)
        self._shape(slide, 0.75, 4.38, 0.9, 0.05, LIME, LIME)
        facts = [_array(value) for value in _array(self._meta.get("cover_facts"))]
        fact_columns = 2
        fact_step = 5.65
        fact_width = 5.15
        for index, fact in enumerate(facts[:COVER_FIELD_LIMITS["cover_facts"]]):
            if len(fact) < 2:
                continue
            column = index % fact_columns
            row = index // fact_columns
            x = 0.75 + column * fact_step
            y = 4.72 + row * 0.68
            self._text_box(slide, x, y, fact_width, 0.2, _text(fact[0]).upper(), 7, MUTED, bold=True)
            self._text_box(slide, x, y + 0.21, fact_width, 0.34, _text(fact[1]), 10, PAPER, bold=True)
        confidentiality = _text(
            self._meta.get("confidentiality"),
            default_confidentiality(self._branding),
        )
        self._text_box(slide, 0.75, 7.0, 8.8, 0.16, confidentiality, 7, MUTED)
        disclosure = _cover_disclosure(self._meta)
        if disclosure:
            self._text_box(slide, 0.75, 7.17, 8.8, 0.14, disclosure, 6.5, MUTED)
        footer_brand = (
            f"Powered by {BRAND_NAME}"
            if client_name and bool(self._branding.get("powered_by", True))
            else ("" if client_name else BRAND_NAME)
        )
        self._text_box(slide, 9.6, 7.08, 3.0, 0.2, footer_brand, 7, MUTED, align="right")

    def _render_sections(self, sections: list[JsonObject]) -> None:
        executive = _object(self._report.get("executive"))
        for section_index, section in enumerate(sections, start=1):
            plans = _plan_section(section)
            if section.get("id") == "executive" and executive and plans:
                lead_blocks = self._executive_blocks(executive)
                first_layout, first_blocks = plans[0]
                if sum(_estimate_block(block) for block in lead_blocks + first_blocks) > 5.15:
                    plans.insert(0, ("text", lead_blocks))
                else:
                    plans[0] = (first_layout, lead_blocks + first_blocks)
            for continuation, (layout, blocks) in enumerate(plans):
                self._render_content_slide(
                    section,
                    blocks,
                    layout,
                    section_index,
                    continuation,
                )

    def _render_deck_slides(self, slides: list[JsonObject]) -> None:
        for index, item in enumerate(slides, start=1):
            plans = _plan_section(item)
            for continuation, (layout, blocks) in enumerate(plans):
                self._render_content_slide(item, blocks, layout, index, continuation)

    def _executive_blocks(self, executive: JsonObject) -> list[JsonObject]:
        blocks: list[JsonObject] = []
        signal = _object(self._meta.get("signal"))
        signal_level = _text(signal.get("level"))
        if signal and signal_level in THEME_SIGNAL:
            signal_color = THEME_SIGNAL[signal_level][1]
            blocks.append(
                {
                    "type": "statement",
                    "text": _text(signal.get("label")) or THEME_SIGNAL[signal_level][0],
                    "tone": "good" if signal_color == THEME_LIME
                    else "watch" if signal_color == THEME_AMBER
                    else "bad" if signal_color == THEME_CORAL
                    else "info",
                }
            )
        bottom_line = _text(executive.get("bottom_line"))
        if bottom_line:
            blocks.append(
                {
                    "type": "callout",
                    "tone": "info",
                    "title": _text(executive.get("label"), "Bottom line"),
                    "text": bottom_line,
                }
            )
        tiles = _array(executive.get("tiles"))
        if tiles:
            blocks.append({"type": "tiles", "tiles": tiles})
        return blocks

    def _render_content_slide(
        self,
        section: JsonObject,
        blocks: list[JsonObject],
        layout: str,
        section_index: int,
        continuation: int,
    ) -> None:
        if layout not in LAYOUTS:
            raise RenderValidationError([f"unsupported inferred layout '{layout}'"])
        if layout in {"cover", "divider"}:
            self._render_divider(section, section_index)
            return
        slide = self._new_slide()
        self._shape(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, PAPER, PAPER)
        title = _text(section.get("title"))
        if continuation:
            title = f"{title} — continued"
        kicker = _text(section.get("kicker"))
        number = _text(section.get("num"), f"{section_index:02d}")
        self._slide_header(slide, title, kicker, number, layout)
        notes: list[str] = []
        if layout == "two-up":
            self._render_two_up_slide(slide, blocks, notes)
        else:
            self._render_block_stack(slide, blocks, 0.7, 1.45, 11.93, 5.35, notes)
        takeaway = _text(section.get("takeaway"))
        if takeaway:
            self._callout(slide, {"tone": "info", "title": "Takeaway", "text": takeaway}, 0.7, 6.35, 11.93, 0.55)
        self._slide_footer(slide, section, blocks)
        if notes:
            self._notes[self._slide_number - 1] = notes

    def _render_divider(self, section: JsonObject, number: int) -> None:
        slide = self._new_slide()
        self._shape(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, INK, INK)
        self._text_box(slide, 0.8, 2.0, 10.8, 0.3, _text(section.get("kicker")).upper(), 9, LIME, bold=True)
        self._text_box(slide, 0.8, 2.46, 10.8, 1.6, _text(section.get("title")), 36, PAPER, bold=True)
        subtitle = _text(section.get("subtitle"))
        if subtitle:
            self._text_box(slide, 0.8, 4.18, 9.7, 0.6, subtitle, 15, SLATE)
        self._text_box(slide, 11.5, 6.9, 1.0, 0.3, f"{number:02d}", 13, LIME, bold=True, align="right")

    def _slide_header(
        self, slide: Any, title: str, kicker: str, number: str, layout: str
    ) -> None:
        accent = LIME if layout != "appendix" else SLATE
        if kicker:
            self._text_box(
                slide,
                0.7,
                0.36,
                10.8,
                0.2,
                kicker.upper(),
                7.5,
                LIME_DARK,
                bold=True,
            )
        self._text_box(slide, 0.7, 0.62, 11.0, 0.58, title, 23, INK, bold=True)
        self._shape(slide, 0.7, 1.22, 11.93, 0.02, INK, INK)
        self._shape(slide, 12.05, 0.55, 0.58, 0.43, accent, accent)
        self._text_box(slide, 12.08, 0.61, 0.52, 0.2, number, 10, INK, bold=True, align="center")

    def _slide_footer(
        self, slide: Any, section: JsonObject, blocks: list[JsonObject]
    ) -> None:
        tags = _source_tags(
            {
                "title": section.get("title"),
                "kicker": section.get("kicker"),
                "blocks": blocks,
            }
        )
        sources = "Sources: " + " · ".join(tags) if tags else "Sources: none cited"
        confidentiality = _text(
            self._meta.get("confidentiality"),
            default_confidentiality(self._branding),
        )
        client_name = _text(self._branding.get("client_name"))
        powered = (
            f" · Powered by {BRAND_NAME}"
            if client_name and bool(self._branding.get("powered_by", True))
            else ""
        )
        data_as_of = _text(self._meta.get("data_as_of"))
        footer_identity = confidentiality
        if data_as_of:
            footer_identity += f" · Data as of {data_as_of}"
        footer_identity += powered
        self._shape(slide, 0.7, 7.05, 11.93, 0.01, RULE, RULE)
        self._text_box(slide, 0.7, 7.11, 4.7, 0.16, sources, 6.5, MUTED)
        self._text_box(
            slide,
            4.9,
            7.11,
            5.8,
            0.16,
            footer_identity,
            6.5,
            MUTED,
            align="center",
        )
        self._text_box(slide, 11.25, 7.11, 1.38, 0.16, str(self._slide_number), 6.5, MUTED, align="right")

    def _render_two_up_slide(
        self, slide: Any, blocks: list[JsonObject], notes: list[str]
    ) -> None:
        if len(blocks) == 1 and _text(blocks[0].get("type")) == "two_col":
            left = [_object(value) for value in _array(blocks[0].get("left"))]
            right = [_object(value) for value in _array(blocks[0].get("right"))]
        else:
            midpoint = max(1, math.ceil(len(blocks) / 2))
            left, right = blocks[:midpoint], blocks[midpoint:]
        self._render_block_stack(slide, left, 0.7, 1.45, 5.73, 5.35, notes)
        self._shape(slide, 6.65, 1.45, 0.01, 5.35, RULE, RULE)
        self._render_block_stack(slide, right, 6.9, 1.45, 5.73, 5.35, notes)

    def _render_block_stack(
        self,
        slide: Any,
        blocks: list[JsonObject],
        x: float,
        y: float,
        width: float,
        height: float,
        notes: list[str],
    ) -> None:
        if not blocks:
            self._text_box(slide, x, y, width, 0.4, "No additional content.", 11, MUTED)
            return
        weights = [max(0.45, _estimate_block(block, width)) for block in blocks]
        usable = height - STACK_GAP * (len(blocks) - 1)
        visual = [_text(block.get("type")) in VISUAL_BLOCKS for block in blocks]
        fixed = sum(weight for weight, is_visual in zip(weights, visual) if not is_visual)
        remainder = usable - fixed
        if any(visual) and remainder >= VISUAL_MIN_HEIGHT * sum(visual):
            # Text keeps its estimated height; visuals share whatever remains.
            heights = [
                remainder / sum(visual) if is_visual else weight
                for weight, is_visual in zip(weights, visual)
            ]
        else:
            total = sum(weights)
            heights = [usable * weight / total for weight in weights]
        cursor = y
        for block, block_height in zip(blocks, heights):
            self._render_block(slide, block, x, cursor, width, block_height, notes)
            cursor += block_height + STACK_GAP

    def _render_block(
        self,
        slide: Any,
        block: JsonObject,
        x: float,
        y: float,
        width: float,
        height: float,
        notes: list[str],
    ) -> None:
        block_type = _text(block.get("type"))
        if block.get("role") == "analysis":
            full = _plain(block.get("text"))
            notes.append(f"{_text(block.get('title'), 'Analysis')}\n{full}")
            summary = deepcopy(block)
            summary["text"] = _compact(full)
            self._callout(slide, summary, x, y, width, height)
            return
        handlers = {
            "text": self._body_text,
            "markdown": self._markdown,
            "bullets": self._bullets,
            "kv": self._kv,
            "table": self._table,
            "tiles": self._tiles,
            "bars": self._bars,
            "pie": self._pie,
            "percentiles": self._percentiles,
            "line": self._line,
            "stacked": self._stacked,
            "waterfall": self._waterfall,
            "band": self._band,
            "heat": self._heat,
            "timeline": self._timeline,
            "callout": self._callout,
            "coverage": self._coverage,
            "findings": self._findings,
            "questions": self._questions,
            "statement": self._statement,
            "chart": self._chart,
        }
        handler = handlers.get(block_type)
        if handler is None:
            self._callout(
                slide,
                {
                    "tone": "watch",
                    "title": "Unsupported block",
                    "text": f"Block type '{block_type}' is not available in PowerPoint.",
                },
                x,
                y,
                width,
                height,
            )
            return
        handler(slide, block, x, y, width, height)

    def _body_text(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        self._block_title(slide, block, x, y, width)
        title_offset = 0.28 if block.get("title") else 0.0
        self._text_box(slide, x, y + title_offset, width, max(0.3, height - title_offset), _plain(block.get("text")), 11, INK_2)

    def _markdown(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        items, paragraphs = _markdown_parts(_text(block.get("text")))
        if items:
            self._bullets(slide, {"type": "bullets", "items": items}, x, y, width, height)
        else:
            self._body_text(slide, {"type": "text", "text": "\n\n".join(paragraphs)}, x, y, width, height)

    def _bullets(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        self._block_title(slide, block, x, y, width)
        offset = 0.3 if block.get("title") else 0.0
        items = [_plain(value) for value in _array(block.get("items"))]
        self._text_box(slide, x, y + offset, width, max(0.3, height - offset), items, 10.5, INK, bullets=True)

    def _kv(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        rows = []
        for value in _array(block.get("rows")):
            row = _array(value)
            if len(row) >= 2:
                suffix = f" [{_text(row[2]).strip('[]')}]" if len(row) > 2 and row[2] else ""
                rows.append([_text(row[0]), _text(row[1]) + suffix])
        table = {"type": "table", "title": block.get("title"), "columns": ["", ""], "rows": rows, "_kv": True}
        self._table(slide, cast(JsonObject, table), x, y, width, height)

    def _table(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        from pptx.util import Inches

        title_offset = self._block_title(slide, block, x, y, width)
        columns = [_text(value) for value in _array(block.get("columns"))]
        rows = [_array(value) for value in _array(block.get("rows"))]
        is_kv = bool(block.get("_kv"))
        if is_kv:
            columns = ["", ""]
        if not columns:
            max_columns = max((len(row) for row in rows), default=0)
            columns = ["" for _ in range(max_columns)]
        table_y = y + title_offset
        table_height = max(0.4, height - title_offset - (0.22 if block.get("note") else 0.0))
        shape = slide.shapes.add_table(
            len(rows) + (0 if is_kv else 1),
            len(columns),
            Inches(x),
            Inches(table_y),
            Inches(width),
            Inches(table_height),
        )
        table = shape.table
        if width > 8 and len(columns) > 1:
            first_width = min(2.2, width * 0.28)
            table.columns[0].width = Inches(first_width)
            remaining_width = (width - first_width) / (len(columns) - 1)
            for column_index in range(1, len(columns)):
                table.columns[column_index].width = Inches(remaining_width)
        row_index = 0
        if not is_kv:
            for column_index, label in enumerate(columns):
                self._cell(table.cell(0, column_index), label, INK, PAPER, True, 7.5)
            row_index = 1
        body_font_size = 6.5 if len(columns) >= 6 else (7.2 if len(columns) >= 5 else 8.0)
        for source_index, row in enumerate(rows):
            for column_index in range(len(columns)):
                value = row[column_index] if column_index < len(row) else ""
                chip_parts = _chip_parts(value)
                if chip_parts is not None:
                    label, raw_status, display_label = chip_parts
                    fill = _status_color(raw_status)
                    detail = _plain(_object(value).get("detail"))
                    self._cell(
                        table.cell(row_index + source_index, column_index),
                        f"● {_status_display(raw_status, label, display_label)}"
                        + (f" — {detail}" if detail else ""),
                        WASH,
                        fill,
                        True,
                        body_font_size,
                    )
                else:
                    fill = PAPER if source_index % 2 == 0 else WASH
                    self._cell(
                        table.cell(row_index + source_index, column_index),
                        _plain(value),
                        fill,
                        INK,
                        column_index == 0,
                        body_font_size,
                    )
        note = _text(block.get("note"))
        if note:
            self._text_box(slide, x, y + height - 0.2, width, 0.18, note, 6.8, MUTED)

    def _tiles(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        title_offset = self._block_title(slide, block, x, y, width)
        tiles = [_object(value) for value in _array(block.get("tiles"))]
        columns = min(4, max(1, len(tiles)))
        rows = math.ceil(len(tiles) / columns)
        gap = 0.12
        tile_width = (width - gap * (columns - 1)) / columns
        tile_height = (height - title_offset - gap * (rows - 1)) / rows
        for index, tile in enumerate(tiles):
            column = index % columns
            row = index // columns
            tx = x + column * (tile_width + gap)
            ty = y + title_offset + row * (tile_height + gap)
            tone = _text(tile.get("tone"))
            color = _tone_color(tone)
            self._shape(slide, tx, ty, tile_width, tile_height, WASH, RULE)
            self._shape(slide, tx, ty, tile_width, 0.05, color, color)
            self._text_box(slide, tx + 0.12, ty + 0.14, tile_width - 0.24, 0.22, _text(tile.get("label")).upper(), 6.5, INK_2, bold=True)
            self._text_box(slide, tx + 0.12, ty + 0.42, tile_width - 0.24, min(0.52, tile_height * 0.35), _text(tile.get("value")), 18, INK, bold=True)
            self._text_box(slide, tx + 0.12, ty + tile_height - 0.36, tile_width - 0.24, 0.25, _plain(tile.get("sub")), 7, INK_2)

    def _bars(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import (
            XL_CHART_TYPE,
            XL_LABEL_POSITION,
            XL_LEGEND_POSITION,
        )
        from pptx.util import Inches, Pt

        title_offset = self._block_title(slide, block, x, y, width)
        series, items, grouped = _normalized_bars(block)
        if _bars_mark(block) == "dots":
            self._paired_dots(
                slide,
                block,
                series,
                items,
                x,
                y + title_offset,
                width,
                max(0.8, height - title_offset),
            )
            return
        data = CategoryChartData()
        data.categories = [_text(item.get("label")) for item in items]
        for series_index, entry in enumerate(series):
            data.add_series(
                _text(entry.get("name"), "Value"),
                [
                    _number(_array(item.get("values"))[series_index])
                    for item in items
                ],
            )
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(x),
            Inches(y + title_offset),
            Inches(width),
            Inches(max(0.8, height - title_offset)),
            data,
        ).chart
        chart.has_legend = grouped
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.font.size = Pt(7)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.tick_labels.font.size = Pt(8)
        chart.category_axis.tick_labels.font.size = Pt(8)
        plot = chart.plots[0]
        plot.gap_width = 70 if grouped else 55
        plot.has_data_labels = True
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.size = Pt(8)
        authored_displays = any(
            bool(item.get("_authored_displays")) for item in items
        )
        for series_index, chart_series in enumerate(plot.series):
            configured = _text(series[series_index].get("color")).lstrip("#")
            chart_series.format.fill.solid()
            chart_series.format.fill.fore_color.rgb = self._rgb(
                configured
                or (
                    SERIES[series_index % len(SERIES)]
                    if grouped
                    else INK
                )
            )
            if authored_displays:
                _set_native_data_label_text(
                    chart_series,
                    [
                        _text(_array(item.get("displays"))[series_index])
                        for item in items
                    ],
                )
            if grouped:
                continue
            for point, item in zip(chart_series.points, items):
                value = _number(_array(item.get("values"))[0])
                point_color = _text(item.get("color")).lstrip("#")
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self._rgb(
                    point_color or (LIME if value >= 0 else CORAL)
                )

    def _paired_dots(
        self,
        slide: Any,
        block: JsonObject,
        series: list[JsonObject],
        items: list[JsonObject],
        x: float,
        y: float,
        width: float,
        height: float,
    ) -> None:
        """Render compact paired dots as editable native PowerPoint shapes."""
        values = [
            _number(value)
            for item in items
            for value in _array(item.get("values"))
        ]
        low, high = min(0.0, min(values)), max(0.0, max(values))
        if block.get("min") is not None:
            low = min(low, _number(block.get("min")))
        if block.get("max") is not None:
            high = max(high, _number(block.get("max")))
        if low == high:
            high = low + 1.0
        label_width = min(1.7, width * 0.28)
        gap_width = min(1.25, width * 0.22)
        plot_x = x + label_width
        plot_width = max(0.8, width - label_width - gap_width)
        legend_height = 0.32
        row_height = max(0.34, (height - legend_height - 0.22) / len(items))
        legend_x = plot_x
        for index, entry in enumerate(series):
            color = _text(entry.get("color")).lstrip("#") or SERIES[
                index % len(SERIES)
            ]
            self._shape(
                slide,
                legend_x,
                y + 0.04,
                0.10,
                0.10,
                color,
                color,
                "oval",
            )
            name = _text(entry.get("name"))
            legend_text_width = min(
                max(0.65, len(name) * 0.065),
                max(0.65, plot_x + plot_width - legend_x),
            )
            self._text_box(
                slide,
                legend_x + 0.14,
                y,
                legend_text_width,
                0.2,
                name,
                7,
                INK_2,
            )
            legend_x += legend_text_width + 0.27
        scale_x = lambda value: (
            plot_x + plot_width * (float(value) - low) / (high - low)
        )
        for item_index, item in enumerate(items):
            row_y = y + legend_height + item_index * row_height
            middle_y = row_y + row_height * 0.45
            self._text_box(
                slide,
                x,
                row_y,
                label_width - 0.10,
                row_height * 0.8,
                _text(item.get("label")),
                7.5,
                INK,
                bold=True,
                align="right",
            )
            self._shape(
                slide,
                plot_x,
                middle_y,
                plot_width,
                0.01,
                RULE,
                RULE,
            )
            item_values = [_number(value) for value in _array(item.get("values"))]
            positions = [scale_x(value) for value in item_values]
            self._shape(
                slide,
                min(positions),
                middle_y - 0.01,
                max(0.01, abs(positions[1] - positions[0])),
                0.02,
                RULE,
                RULE,
            )
            for series_index, position in enumerate(positions):
                color = (
                    _text(series[series_index].get("color")).lstrip("#")
                    or SERIES[series_index % len(SERIES)]
                )
                self._shape(
                    slide,
                    position - 0.055,
                    middle_y - 0.055,
                    0.11,
                    0.11,
                    color,
                    PAPER,
                    "oval",
                )
            self._text_box(
                slide,
                plot_x + plot_width + 0.08,
                row_y,
                gap_width - 0.08,
                row_height * 0.8,
                _bar_gap_display(
                    item,
                    item_values[1] - item_values[0],
                    block,
                ),
                7,
                INK,
                bold=True,
            )
        axis_y = y + legend_height + len(items) * row_height
        self._text_box(
            slide,
            plot_x,
            axis_y,
            0.6,
            0.18,
            f"{low:g}",
            7,
            MUTED,
        )
        self._text_box(
            slide,
            plot_x + plot_width - 0.6,
            axis_y,
            0.6,
            0.18,
            f"{high:g}",
            7,
            MUTED,
            align="right",
        )

    def _pie(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt

        title_offset = self._block_title(slide, block, x, y, width)
        items = [
            item
            for value in _array(block.get("items"))
            if (item := _object(value)) and _number(item.get("value")) > 0
        ]
        data = CategoryChartData()
        data.categories = [_text(item.get("label")) for item in items]
        data.add_series("Value", [_number(item.get("value")) for item in items])
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            Inches(x),
            Inches(y + title_offset),
            Inches(width),
            Inches(max(0.8, height - title_offset)),
            data,
        ).chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        plot = chart.plots[0]
        plot.hole_size = 58
        plot.has_data_labels = True
        plot.data_labels.position = XL_LABEL_POSITION.BEST_FIT
        plot.data_labels.show_value = True
        plot.data_labels.font.size = Pt(8)
        series = plot.series[0]
        for index, (point, item) in enumerate(zip(series.points, items)):
            configured = _text(item.get("color")).lstrip("#")
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = self._rgb(configured or PIE[index])

    def _percentiles(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        items: list[JsonValue] = []
        for value in _array(block.get("items")):
            item = _object(value)
            percentile = _number(item.get("percentile"))
            items.append(
                {
                    "label": item.get("label", ""),
                    "value": percentile,
                    "display": item.get("value_display") or _ordinal(percentile),
                }
            )
        self._bars(
            slide,
            {
                "type": "bars",
                "title": block.get("title"),
                "items": items,
                "max": 100,
                "note": block.get("note"),
            },
            x,
            y,
            width,
            height,
        )

    def _line(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt

        title_offset = self._block_title(slide, block, x, y, width)
        series = [_object(value) for value in _array(block.get("series"))]
        categories: list[str] = []
        for item in series:
            points = [_array(value) for value in _array(item.get("points"))]
            if points:
                categories = [_text(point[0]) for point in points if point]
                break
        data = CategoryChartData()
        data.categories = categories
        for item in series:
            point_map = {
                _text(point[0]): _number(point[1])
                for value in _array(item.get("points"))
                if len(point := _array(value)) >= 2
            }
            data.add_series(
                _text(item.get("name"), "Series"),
                [point_map.get(category, None) for category in categories],
            )
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE_MARKERS,
            Inches(x),
            Inches(y + title_offset),
            Inches(width),
            Inches(max(0.8, height - title_offset)),
            data,
        ).chart
        narrow = width < 7.0
        chart.has_legend = len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.font.size = Pt(7)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.tick_labels.font.size = Pt(8)
        values = [
            _number(point[1])
            for item in series
            for value in _array(item.get("points"))
            if len(point := _array(value)) >= 2
        ]
        if values:
            if narrow:
                low, high = min(values), max(values)
                if low == high:
                    padding = abs(low) * 0.1 or 1.0
                    low, high = low - padding, high + padding
                raw_step = (high - low) / 4
                magnitude = 10 ** math.floor(math.log10(raw_step))
                fraction = raw_step / magnitude
                nice_fraction = next(
                    candidate
                    for candidate in (1.0, 2.0, 2.5, 5.0, 10.0)
                    if fraction <= candidate
                )
                major_unit = nice_fraction * magnitude
                chart.value_axis.minimum_scale = (
                    math.floor(low / major_unit) * major_unit
                )
                chart.value_axis.maximum_scale = (
                    math.ceil(high / major_unit) * major_unit
                )
                chart.value_axis.major_unit = major_unit
            chart.value_axis.tick_labels.number_format = _axis_number_format(
                max(values, key=abs),
                decimals=int(_number(block.get("decimals"), 1)),
                currency=_text(block.get("currency")),
                value_format=_text(block.get("y_format")),
                suffix=_text(block.get("y_suffix")),
            )
            chart.value_axis.tick_labels.number_format_is_linked = False
        chart.category_axis.tick_labels.font.size = Pt(7)
        for index, chart_series in enumerate(chart.series):
            chart_series.format.line.color.rgb = self._rgb(SERIES[index % len(SERIES)])
            chart_series.format.line.width = Pt(1.8)

    def _stacked(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt

        title_offset = self._block_title(slide, block, x, y, width)
        raw_items = [_object(value) for value in _array(block.get("items"))]
        rows = (
            raw_items
            if raw_items and _array(raw_items[0].get("segments"))
            else [{"label": _text(block.get("label"), "Mix"), "segments": raw_items}]
        )
        categories = [_text(row.get("label"), "Mix") for row in rows]
        segment_names: list[str] = []
        for row in rows:
            for value in _array(row.get("segments")):
                name = _text(_object(value).get("label"))
                if name and name not in segment_names:
                    segment_names.append(name)
        data = CategoryChartData()
        data.categories = categories
        for name in segment_names:
            amounts: list[float] = []
            for row in rows:
                by_name = {
                    _text(segment.get("label")): _number(segment.get("value"))
                    for value in _array(row.get("segments"))
                    if (segment := _object(value))
                }
                amounts.append(by_name.get(name, 0.0))
            data.add_series(name, amounts)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_STACKED,
            Inches(x),
            Inches(y + title_offset),
            Inches(width),
            Inches(max(0.8, height - title_offset)),
            data,
        ).chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.value_axis.has_major_gridlines = True
        chart.category_axis.tick_labels.font.size = Pt(8)
        for index, chart_series in enumerate(chart.series):
            chart_series.format.fill.solid()
            chart_series.format.fill.fore_color.rgb = self._rgb(SERIES[index % len(SERIES)])

    def _waterfall(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        title_offset = self._block_title(slide, block, x, y, width)
        items = [_object(value) for value in _array(block.get("items"))]
        bases: list[float] = []
        increases: list[float] = []
        decreases: list[float] = []
        running = 0.0
        for item in items:
            amount = _number(item.get("value"))
            is_total = bool(item.get("total") or item.get("is_total"))
            if is_total:
                total = amount if item.get("value") is not None else running
                bases.append(0.0)
                increases.append(max(0.0, total))
                decreases.append(max(0.0, -total))
                running = total
            elif amount >= 0:
                bases.append(max(0.0, running))
                increases.append(amount)
                decreases.append(0.0)
                running += amount
            else:
                running += amount
                bases.append(max(0.0, running))
                increases.append(0.0)
                decreases.append(abs(amount))
        data = CategoryChartData()
        data.categories = [_text(item.get("label")) for item in items]
        data.add_series("Base", bases)
        data.add_series("Increase", increases)
        data.add_series("Decrease", decreases)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_STACKED,
            Inches(x),
            Inches(y + title_offset),
            Inches(width),
            Inches(max(0.8, height - title_offset)),
            data,
        ).chart
        chart.has_legend = False
        chart.value_axis.has_major_gridlines = True
        chart.series[0].format.fill.background()
        chart.series[0].format.line.fill.background()
        chart.series[1].format.fill.solid()
        chart.series[1].format.fill.fore_color.rgb = self._rgb(LIME)
        chart.series[2].format.fill.solid()
        chart.series[2].format.fill.fore_color.rgb = self._rgb(CORAL)

    def _band(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        items = [_object(value) for value in _array(block.get("items"))]
        title_offset = self._block_title(slide, block, x, y, width)
        values = [
            _number(value)
            for item in items
            for value in (
                item.get("low", item.get("min")),
                item.get("high", item.get("max")),
                item.get("value", item.get("current")),
                item.get("target", item.get("value", item.get("current"))),
            )
        ]
        low, high = min(values, default=0.0), max(values, default=1.0)
        span = (high - low) or 1.0
        label_width = min(2.1, width * 0.3)
        plot_x = x + label_width
        plot_width = max(0.8, width - label_width - 0.35)
        row_height = max(0.32, (height - title_offset) / max(len(items), 1))
        for index, item in enumerate(items):
            row_y = y + title_offset + index * row_height
            self._text_box(
                slide,
                x,
                row_y + 0.04,
                label_width - 0.1,
                row_height - 0.05,
                _text(item.get("label")),
                8,
                INK,
                bold=True,
                align="right",
            )
            item_low = _number(item.get("low", item.get("min")))
            item_high = _number(item.get("high", item.get("max")))
            current = _number(item.get("value", item.get("current")))
            target = _number(item.get("target"), current)
            range_x = plot_x + plot_width * (item_low - low) / span
            range_width = max(0.03, plot_width * (item_high - item_low) / span)
            self._shape(
                slide,
                range_x,
                row_y + row_height * 0.38,
                range_width,
                0.11,
                RULE,
                RULE,
            )
            current_x = plot_x + plot_width * (current - low) / span
            target_x = plot_x + plot_width * (target - low) / span
            self._shape(
                slide,
                current_x - 0.04,
                row_y + row_height * 0.3,
                0.08,
                0.24,
                LIME_DARK,
                LIME_DARK,
                "oval",
            )
            self._shape(
                slide,
                target_x - 0.015,
                row_y + row_height * 0.2,
                0.03,
                0.38,
                INK,
                INK,
            )

    def _heat(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        from pptx.util import Inches

        title_offset = self._block_title(slide, block, x, y, width)
        columns = [_text(value) for value in _array(block.get("columns"))]
        raw_rows = _array(block.get("rows"))
        object_rows = [_object(value) for value in raw_rows]
        if object_rows and object_rows[0]:
            row_labels = [_text(row.get("label")) for row in object_rows]
            values = [_array(row.get("values")) for row in object_rows]
        else:
            array_rows = [_array(value) for value in raw_rows]
            row_labels = [_text(row[0]) for row in array_rows if row]
            values = [row[1:] for row in array_rows if row]
        if len(columns) == (len(values[0]) + 1 if values else 1):
            columns = columns[1:]
        if not columns and values:
            columns = [str(index + 1) for index in range(len(values[0]))]
        shape = slide.shapes.add_table(
            len(values) + 1,
            len(columns) + 1,
            Inches(x),
            Inches(y + title_offset),
            Inches(width),
            Inches(max(0.8, height - title_offset)),
        )
        table = shape.table
        self._cell(table.cell(0, 0), "", INK, PAPER, True, 7)
        for index, label in enumerate(columns, start=1):
            self._cell(table.cell(0, index), label, INK, PAPER, True, 7)
        numeric = [_number(value) for row in values for value in row]
        low = min(numeric, default=0.0)
        high = max(numeric, default=1.0)
        for row_index, row in enumerate(values, start=1):
            label = row_labels[row_index - 1] if row_index - 1 < len(row_labels) else str(row_index)
            self._cell(table.cell(row_index, 0), label, INK, PAPER, True, 7.5)
            for column_index in range(1, len(columns) + 1):
                raw = row[column_index - 1] if column_index - 1 < len(row) else None
                value = _number(raw)
                ratio = (value - low) / ((high - low) or 1.0)
                color = self._mix_hex(WASH, LIME, ratio)
                self._cell(table.cell(row_index, column_index), _text(raw, "—"), color, INK, False, 7.5)

    def _timeline(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        items = [_object(value) for value in _array(block.get("items"))]
        offset = self._block_title(slide, block, x, y, width)
        top = y + offset + 0.08
        usable = max(0.8, height - offset - 0.12)
        row_height = usable / max(1, len(items))
        line_x = x + min(1.35, width * 0.22)
        self._shape(slide, line_x, top, 0.025, usable, RULE, RULE)
        for index, item in enumerate(items):
            row_y = top + index * row_height
            tone = _text(item.get("tone"), "info")
            color = _tone_color(tone)
            marker_y = row_y + min(0.18, row_height * 0.25)
            self._shape(
                slide,
                line_x - 0.055,
                marker_y,
                0.14,
                0.14,
                color,
                PAPER,
                "oval",
            )
            self._text_box(
                slide,
                x,
                row_y,
                max(0.7, line_x - x - 0.14),
                min(0.3, row_height),
                _text(item.get("date")),
                7.5,
                INK_2,
                bold=True,
                align="right",
            )
            content_x = line_x + 0.2
            self._text_box(
                slide,
                content_x,
                row_y,
                max(0.8, x + width - content_x),
                min(0.28, row_height * 0.42),
                _plain(item.get("title")),
                9.5,
                INK,
                bold=True,
            )
            detail = _plain(item.get("detail"))
            if detail:
                self._text_box(
                    slide,
                    content_x,
                    row_y + min(0.29, row_height * 0.42),
                    max(0.8, x + width - content_x),
                    max(0.24, row_height * 0.52),
                    detail,
                    8.5,
                    INK_2,
                )

    def _callout(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        color = _tone_color(_text(block.get("tone"), "info"))
        self._shape(slide, x, y, width, height, WASH, RULE)
        self._shape(slide, x, y, 0.06, height, color, color)
        title = _text(block.get("title"))
        if title:
            self._text_box(slide, x + 0.18, y + 0.12, width - 0.32, 0.25, title, 10, INK, bold=True)
        text_y = y + (0.43 if title else 0.15)
        self._text_box(slide, x + 0.18, text_y, width - 0.32, max(0.25, height - (text_y - y) - 0.1), _plain(without_lead(block.get("text"), title)), 9.5, INK_2)

    def _coverage(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        rows: list[JsonValue] = []
        for value in _array(block.get("items")):
            item = _object(value)
            raw_status = _text(item.get("status"))
            rows.append(
                [
                    item.get("name", ""),
                    {
                        "chip": _text(item.get("label")) or raw_status,
                        "status": raw_status,
                        "display_label": _text(item.get("display_label")),
                    },
                    without_lead(
                        item.get("note", ""),
                        status_label(raw_status, item.get("display_label")),
                        item.get("label"),
                    ),
                ]
            )
        self._table(
            slide,
            {"type": "table", "title": block.get("title"), "columns": ["Source", "Status", "Note"], "rows": rows},
            x,
            y,
            width,
            height,
        )

    def _findings(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        items = [_object(value) for value in _array(block.get("items"))]
        if not items:
            self._callout(
                slide,
                {
                    "type": "callout",
                    "tone": "good",
                    "title": block.get("empty_title", "No findings"),
                    "text": block.get("empty_text", ""),
                },
                x,
                y,
                width,
                height,
            )
            return
        rows: list[JsonValue] = []
        for item in items:
            severity = _text(item.get("severity"), "low")
            rows.append(
                [
                    {"chip": severity.upper(), "status": severity},
                    item.get("title", ""),
                    item.get("detail", ""),
                ]
            )
        self._table(
            slide,
            {"type": "table", "title": block.get("title"), "columns": ["Severity", "Finding", "Detail"], "rows": rows},
            x,
            y,
            width,
            height,
        )

    def _questions(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        items = _array(block.get("items"))
        questions: list[str] = []
        for index, value in enumerate(items, start=1):
            item = _object(value)
            question = _text(item.get("q") if item else value)
            why = _text(item.get("why"))
            questions.append(f"{index:02d}  {question}" + (f"\n      {why}" if why else ""))
        self._text_box(slide, x, y, width, height, questions, 10, INK, bullets=False)

    def _statement(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        color = _tone_color(_text(block.get("tone"), "good"))
        self._shape(slide, x, y + 0.08, 0.08, max(0.5, height - 0.16), color, color)
        self._text_box(slide, x + 0.28, y + 0.05, width - 0.35, min(1.1, height * 0.58), _plain(block.get("text")), 18, INK, bold=True)
        sub = _plain(block.get("sub"))
        if sub:
            self._text_box(slide, x + 0.28, y + min(1.2, height * 0.6), width - 0.35, max(0.28, height * 0.3), sub, 9, INK_2)

    def _chart(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float, height: float
    ) -> None:
        chart = _object(block.get("chart"))
        status = _text(chart.get("status"))
        if status != "ok":
            self._callout(
                slide,
                {
                    "type": "callout",
                    "tone": "info",
                    "title": chart.get("title", "Chart unavailable"),
                    "text": chart.get("error") or status,
                },
                x,
                y,
                width,
                height,
            )
            return
        hint = _object(chart.get("render_hint"))
        kind = effective_chart_mode(chart)
        columns = [_object(value) for value in _array(chart.get("columns"))]
        keys = [_text(column.get("key")) for column in columns]
        rows = [_array(value) for value in _array(chart.get("rows"))]
        x_key = _text(hint.get("x"))
        y_keys = [_text(value) for value in _array(hint.get("y"))]
        if kind == "line" and x_key in keys:
            x_index = keys.index(x_key)
            series: list[JsonValue] = []
            for y_key in y_keys:
                if y_key not in keys:
                    continue
                y_index = keys.index(y_key)
                series.append(
                    {
                        "name": _text(columns[y_index].get("title") or y_key),
                        "points": [
                            [row[x_index], self._chart_value(columns[y_index], row[y_index])]
                            for row in rows
                            if len(row) > max(x_index, y_index)
                        ],
                    }
                )
            self._line(
                slide,
                {
                    "type": "line",
                    "title": chart.get("title"),
                    "series": series,
                    "y_suffix": "%" if (
                        y_keys
                        and y_keys[0] in keys
                        and _text(columns[keys.index(y_keys[0])].get("format"))
                        == "percentage"
                    ) else "",
                    "decimals": (
                        columns[keys.index(y_keys[0])].get("decimals")
                        if y_keys and y_keys[0] in keys
                        else 1
                    ),
                    "y_format": (
                        columns[keys.index(y_keys[0])].get("format")
                        if y_keys and y_keys[0] in keys
                        else ""
                    ),
                    "currency": chart.get("currency"),
                },
                x,
                y,
                width,
                height,
            )
            return
        unit_index = keys.index("unit") if "unit" in keys else None
        mixed_units = unit_index is not None and len({
            _text(row[unit_index])
            for row in rows
            if len(row) > unit_index
        }) > 1
        if kind in {"bars", "pie"} and not mixed_units and x_key in keys and y_keys and y_keys[0] in keys:
            x_index = keys.index(x_key)
            y_index = keys.index(y_keys[0])
            items: list[JsonValue] = [
                {
                    "label": row[x_index],
                    "value": (
                        _number(row[y_index])
                        if kind == "pie"
                        else self._chart_value(columns[y_index], row[y_index])
                    ),
                }
                for row in rows
                if len(row) > max(x_index, y_index)
            ]
            item_values = [_number(_object(item).get("value")) for item in items]
            if kind == "pie" and (
                any(value < 0 for value in item_values)
                or not any(value > 0 for value in item_values)
            ):
                items = []
            elif kind == "pie" and sum(value > 0 for value in item_values) > len(PIE):
                kind = "bars"
            handler = self._pie if kind == "pie" else self._bars
            if items:
                handler(
                    slide,
                    {"type": kind, "title": chart.get("title"), "items": items},
                    x,
                    y,
                    width,
                    height,
                )
                return
        formatted_rows: list[JsonValue] = []
        for row in rows:
            formatted_rows.append(
                [
                    self._format_chart_cell(column, row[index] if index < len(row) else None, _text(chart.get("currency")))
                    for index, column in enumerate(columns)
                ]
            )
        note = ""
        if chart.get("_appendix_reference"):
            note = (
                f"Showing {len(formatted_rows)} of "
                f"{int(_number(chart.get('_preview_total')))} returned rows; "
                f"all returned rows are in "
                f"{_text(chart.get('_appendix_reference'))}"
            )
            if chart.get("_source_truncated"):
                note += ". Connector output was truncated at source"
        elif chart.get("_chart_data_appendix") and chart.get("truncated"):
            note = (
                "Connector output was truncated at source; this appendix contains "
                "all returned rows, not the complete dataset"
            )
        elif chart.get("truncated"):
            note = "Connector output was truncated at source"
        self._table(
            slide,
            {
                "type": "table",
                "title": chart.get("title"),
                "columns": [column.get("title") or column.get("key") or "" for column in columns],
                "rows": formatted_rows[:TABLE_ROWS_PER_SLIDE],
                "note": note,
            },
            x,
            y,
            width,
            height,
        )

    def _chart_value(self, column: JsonObject, value: JsonValue) -> float:
        amount = _number(value)
        return amount * 100.0 if _text(column.get("format")) == "percentage" else amount

    def _format_chart_cell(
        self, column: JsonObject, value: JsonValue, currency: str
    ) -> str:
        if value is None:
            return "—"
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            return _text(value)
        decimals = int(_number(column.get("decimals"), 2))
        kind = _text(column.get("format"))
        if kind == "percentage":
            return f"{float(value) * 100:.{decimals}f}%"
        if kind == "currency":
            prefix = f"{currency} " if currency else ""
            return f"{prefix}{float(value):,.{decimals}f}"
        if kind in {"number", "decimal"}:
            return f"{float(value):,.{decimals}f}"
        return _text(value)

    def _block_title(
        self, slide: Any, block: JsonObject, x: float, y: float, width: float
    ) -> float:
        title = _text(block.get("title"))
        if not title:
            return 0.0
        self._text_box(slide, x, y, width, 0.22, title.upper(), 7.2, INK_2, bold=True)
        return 0.3

    def _write_notes(self) -> None:
        for slide_index, entries in self._notes.items():
            slide = self._presentation.slides[slide_index]
            try:
                frame = slide.notes_slide.notes_text_frame
                existing = frame.text.strip()
                payload = "\n\n".join(entries)
                frame.text = f"{existing}\n\n{payload}".strip()
            except (AttributeError, NotImplementedError):
                continue

    def _validate_presentation(self) -> list[str]:
        errors: list[str] = list(self._layout_errors)
        slide_width = int(self._presentation.slide_width)
        slide_height = int(self._presentation.slide_height)
        footer_limit = int(slide_height * 0.947)
        for slide_index, slide in enumerate(self._presentation.slides, start=1):
            visible_text = ""
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    visible_text += "\n" + shape.text
                is_background = (
                    shape.left == 0
                    and shape.top == 0
                    and shape.width == slide_width
                    and shape.height == slide_height
                )
                is_footer = shape.top >= footer_limit
                if shape.left < 0 or shape.top < 0:
                    if slide_index != 1:
                        errors.append(f"slide {slide_index} has a shape outside the top/left boundary")
                if shape.left + shape.width > slide_width + 1:
                    errors.append(f"slide {slide_index} has horizontal overflow")
                if shape.top + shape.height > slide_height + 1:
                    errors.append(f"slide {slide_index} has vertical overflow")
                if (
                    slide_index > 1
                    and not is_background
                    and not is_footer
                    and shape.top < footer_limit < shape.top + shape.height
                ):
                    errors.append(f"slide {slide_index} content overlaps the footer")
            if not visible_text.strip():
                errors.append(f"slide {slide_index} is an empty placeholder")
        return list(dict.fromkeys(errors))

    def _shape(
        self,
        slide: Any,
        x: float,
        y: float,
        width: float,
        height: float,
        fill: str,
        line: str,
        shape_kind: str = "rect",
    ) -> Any:
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches

        shape_type = MSO_SHAPE.OVAL if shape_kind == "oval" else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(
            shape_type, Inches(x), Inches(y), Inches(width), Inches(height)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb(fill)
        shape.line.color.rgb = self._rgb(line)
        return shape

    def _text_box(
        self,
        slide: Any,
        x: float,
        y: float,
        width: float,
        height: float,
        content: str | list[str],
        size: float,
        color: str,
        *,
        bold: bool = False,
        align: str = "left",
        bullets: bool = False,
    ) -> Any:
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches, Pt

        box = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(width), Inches(height)
        )
        frame = box.text_frame
        frame.clear()
        frame.margin_left = Inches(0.01)
        frame.margin_right = Inches(0.01)
        frame.margin_top = Inches(0.01)
        frame.margin_bottom = Inches(0.01)
        frame.vertical_anchor = MSO_ANCHOR.TOP
        values = content if isinstance(content, list) else [content]
        characters_per_line = max(1, int(width * 72 / max(size * 0.52, 1)))
        estimated_lines = sum(
            max(1, math.ceil(len(str(value)) / characters_per_line))
            for value in values
        )
        available_lines = max(1, int(height * 72 / max(size * 1.2, 1)))
        if estimated_lines > available_lines + 1:
            self._layout_errors.append(
                f"slide {self._slide_number} text overflows its frame "
                f"({estimated_lines} estimated lines, {available_lines} available)"
            )
        for index, value in enumerate(values):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = value
            paragraph.font.name = FONT_FAMILY
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = self._rgb(color)
            paragraph.alignment = {
                "center": PP_ALIGN.CENTER,
                "right": PP_ALIGN.RIGHT,
            }.get(align, PP_ALIGN.LEFT)
            paragraph.space_after = Pt(5 if len(values) > 1 else 0)
            if bullets:
                paragraph.level = 0
                paragraph.text = f"•  {value}"
        return box

    def _cell(
        self,
        cell: Any,
        text: str,
        fill: str,
        color: str,
        bold: bool,
        size: float,
    ) -> None:
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Inches, Pt

        cell.fill.solid()
        cell.fill.fore_color.rgb = self._rgb(fill)
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame = cell.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text = text
        paragraph.font.name = FONT_FAMILY
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = self._rgb(color)

    def _rgb(self, value: str) -> Any:
        from pptx.dml.color import RGBColor

        normalized = value.lstrip("#")
        return RGBColor(
            int(normalized[0:2], 16),
            int(normalized[2:4], 16),
            int(normalized[4:6], 16),
        )

    def _mix_hex(self, start: str, end: str, ratio: float) -> str:
        clamped = min(1.0, max(0.0, ratio))
        start_rgb = [int(start[index : index + 2], 16) for index in (0, 2, 4)]
        end_rgb = [int(end[index : index + 2], 16) for index in (0, 2, 4)]
        mixed = [
            round(left + (right - left) * clamped)
            for left, right in zip(start_rgb, end_rgb)
        ]
        return "".join(f"{component:02X}" for component in mixed)


def render_report(
    report: JsonObject,
    output_path: str | Path,
    template_path: str | Path | None = None,
    branding: JsonObject | None = None,
) -> None:
    """Render one report object to an editable PowerPoint presentation."""
    resolved_template = Path(template_path) if template_path is not None else None
    if resolved_template is None:
        here = Path(__file__).resolve().parent
        resolved_template = next(
            (
                candidate
                for candidate in (
                    here.parent / "assets" / "gradient-master.potx",
                    here.parent.parent.parent / "assets" / "gradient-master.potx",
                )
                if candidate.is_file()
            ),
            None,
        )
    renderer = PptxRenderer(report, resolved_template, branding)
    renderer.render(Path(output_path))


def load_report(path: str | Path) -> JsonObject:
    """Read a report JSON object with stable validation errors."""
    report_path = Path(path)
    try:
        value = cast(JsonValue, json.loads(report_path.read_text(encoding="utf-8")))
    except OSError as error:
        raise RenderValidationError([f"unable to read {report_path}: {error}"]) from error
    except json.JSONDecodeError as error:
        raise RenderValidationError(
            [
                f"invalid JSON in {report_path} at line {error.lineno}, "
                f"column {error.colno}: {error.msg}"
            ]
        ) from error
    if not isinstance(value, dict):
        raise RenderValidationError(["report root must be a JSON object"])
    return cast(JsonObject, value)
