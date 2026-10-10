#!/usr/bin/env python3
"""Focused release-quality checks for the shared report toolchain."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from collections.abc import Callable
from typing import Any


ROOT = pathlib.Path(__file__).resolve().parents[1]
SHARED = ROOT / "shared"
FIXTURES = ROOT / "tests" / "fixtures"
GOLDEN_REPORT = FIXTURES / "golden-report.json"
UNIFIED_RENDERER = SHARED / "render.py"
GOLDEN_GENERATOR = ROOT / "tools" / "generate_golden_example.py"

Check = Callable[[bool, str], None]


def _load_module(name: str, path: pathlib.Path) -> Any:
    shared = str(SHARED)
    if shared not in sys.path:
        sys.path.insert(0, shared)
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _slide_text(slide: Any) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def _outline_titles(items: list[Any]) -> list[str]:
    titles: list[str] = []
    for item in items:
        if isinstance(item, list):
            titles.extend(_outline_titles(item))
        else:
            title = getattr(item, "title", None)
            if title:
                titles.append(str(title))
    return titles


def _pdf_base_fonts(reader: Any) -> set[str]:
    fonts: set[str] = set()
    for page in reader.pages:
        resources = page.get("/Resources")
        if resources is None:
            continue
        resources = resources.get_object()
        font_map = resources.get("/Font")
        if font_map is None:
            continue
        for reference in font_map.get_object().values():
            font = reference.get_object()
            base_font = font.get("/BaseFont")
            if base_font:
                fonts.add(str(base_font))
    return fonts


def run_source_checks(check: Check) -> None:
    """Exercise pure rendering, composition and layout contracts."""
    report = _load_module("release_gradient_report", SHARED / "gradient_report.py")
    pptx = _load_module("release_gradient_pptx", SHARED / "gradient_pptx.py")
    composer = _load_module("release_compose_report", SHARED / "compose_report.py")
    layout = _load_module("release_check_layout", SHARED / "check_layout.py")

    check(
        "font-family:Inter" in report.CSS
        and 'paragraph.font.name = "Inter"' in (SHARED / "gradient_pptx.py").read_text(encoding="utf-8")
        and (ROOT / "assets" / "fonts" / "OFL.txt").is_file()
        and all(
            (ROOT / "assets" / "fonts" / filename).is_file()
            for filename in (
                "inter-latin-wght-normal.woff2",
                "inter-latin-ext-wght-normal.woff2",
            )
        ),
        "shared PDF and PowerPoint renderers require licensed Inter typography",
    )

    signed = report.b_bars(
        {
            "title": "Signed bars (TEST)",
            "items": [
                {"label": "Positive (TEST)", "value": 2.0, "display": "+2.0%"},
                {"label": "Negative (TEST)", "value": -1.0, "display": "-1.0%"},
            ],
        }
    )
    check(
        signed.count("<rect") >= 4
        and report.LIME in signed
        and report.CORAL in signed
        and "+2.0%" in signed
        and "-1.0%" in signed,
        "signed bars render positive and negative values around zero",
    )

    new_blocks = {
        "pie": report.b_pie(
            {
                "title": "Pie (TEST)",
                "items": [
                    {"label": "Public (TEST)", "value": 60, "display": "60%"},
                    {"label": "Private (TEST)", "value": 40, "display": "40%"},
                ],
                "center": "100%",
            }
        ),
        "stacked": report.b_stacked(
            {
                "title": "Stacked (TEST)",
                "items": [
                    {"label": "A (TEST)", "value": 60},
                    {"label": "B (TEST)", "value": 40},
                ],
            }
        ),
        "waterfall": report.b_waterfall(
            {
                "title": "Waterfall (TEST)",
                "items": [
                    {"label": "Start", "value": 10, "total": True},
                    {"label": "Loss", "value": -2},
                ],
            }
        ),
        "band": report.b_band(
            {
                "title": "Band (TEST)",
                "items": [
                    {
                        "label": "Range (TEST)",
                        "min": 1,
                        "max": 4,
                        "target": 2,
                        "current": 3,
                    }
                ],
            }
        ),
        "heat": report.b_heat(
            {
                "title": "Heat (TEST)",
                "columns": ["Metric", "A"],
                "rows": [["Return", -0.2]],
            }
        ),
        "timeline": report.b_timeline(
            {
                "title": "Timeline (TEST)",
                "items": [
                    {
                        "date": "2026-10-01",
                        "title": "Example event (TEST)",
                        "detail": "Fictional event detail.",
                        "tone": "watch",
                    }
                ],
            }
        ),
    }
    check(
        all(
            "<svg" in value
            for key, value in new_blocks.items()
            if key not in {"heat", "timeline"}
        )
        and "<table" in new_blocks["heat"]
        and "timeline-row" in new_blocks["timeline"],
        "pie, stacked, waterfall, band, heat, and timeline blocks render",
    )

    markdown = (
        "# Flow test\n\nPreamble.\n\n"
        "## 1. First\n\nFirst body.\n\n"
        "## 2. Second\n\nSecond body.\n\n"
        "## Appendix A — Sources\n\nSource body."
    )
    flowed = report.md_document_to_report(markdown, {})
    sections = flowed["sections"]
    check(
        len(sections) == 3
        and sections[0]["new_page"] is False
        and sections[1]["new_page"] is False
        and sections[2]["new_page"] is True,
        "markdown sections flow by default and appendices start a page",
    )
    body = report.body_html(flowed)
    check(
        body.count('class="sec new-page"') == 1
        and '<div class="keep"><div class="sh">' in body,
        "section headings stay with opening content",
    )
    check(
        "thead{display:table-header-group}" in report.CSS
        and "tr{break-inside:avoid}" in report.CSS,
        "PDF tables repeat headers and retain rows across page splits",
    )

    table = {
        "type": "table",
        "title": "Long table (TEST)",
        "columns": ["Item", "Value"],
        "rows": [[f"Row {index:02d} (TEST)", index] for index in range(25)],
    }
    chunks = pptx._split_table_block(table)
    check(
        [len(chunk["rows"]) for chunk in chunks] == [12, 12, 1]
        and chunks[0]["title"] == "Long table (TEST)"
        and all("continued" in chunk["title"] for chunk in chunks[1:]),
        "PowerPoint tables split deterministically with retained headings",
    )

    composed = composer.compose(
        "# Composition (TEST)\n\n## 1. Summary\n\nAuthoritative body [S1].",
        {
            "sections": {
                "1": {
                    "before": [
                        {
                            "type": "unavailable",
                            "title": "Optional evidence",
                            "reason": "fictional source omitted",
                            "source_tags": ["S1"],
                        }
                    ],
                    "analysis": [
                        {
                            "title": "test conclusion",
                            "text": "A fictional analytical conclusion [S1].",
                        }
                    ],
                }
            }
        },
        {"title": "Composition (TEST)"},
    )
    blocks = composed["sections"][0]["blocks"]
    check(
        [block["role"] for block in blocks if "role" in block]
        == ["unavailable", "analysis"]
        and any(
            block.get("type") == "markdown"
            and block.get("text") == "Authoritative body [S1]."
            for block in blocks
        ),
        "generic composition preserves markdown and typed placements",
    )

    original_bbox = layout.bbox_pages
    original_lines = layout.page_lines
    try:
        layout.bbox_pages = lambda _path: [
            (100.0, [(10.0, 70.0)]),
            (100.0, [(10.0, 25.0)]),
        ]
        layout.page_lines = lambda _path: [
            ["Body", "continues", "normally"],
            ["Section heading", "Footer"],
        ]
        failures = layout.check_layout(
            pathlib.Path("fictional.pdf"),
            headings=["Section heading"],
            minimum_fill=0.40,
            excluded=set(),
            minimum_pages=3,
            maximum_pages=4,
        )
    finally:
        layout.bbox_pages = original_bbox
        layout.page_lines = original_lines
    check(
        any("page count 2 is below 3" in failure for failure in failures)
        and any("page 2 is 15% full" in failure for failure in failures)
        and any("ends with heading" in failure for failure in failures)
        and layout.parse_pages("1,3-4") == {1, 3, 4},
        "layout checker catches page count, sparse pages, and orphan headings",
    )

    golden = json.loads(GOLDEN_REPORT.read_text(encoding="utf-8"))
    serialized = json.dumps(golden, ensure_ascii=False)
    check(
        golden["meta"]["title"].endswith("(TEST)")
        and "Fictional" in serialized
        and not re.search(
            r"(?i)(api[_-]?key|bearer\s+[a-z0-9._-]+|password\s*[:=]|secret\s*[:=])",
            serialized,
        ),
        "golden source is explicitly fictional and contains no credentials",
    )


def _check_pdf_quality(check: Check, pdf_path: pathlib.Path) -> None:
    try:
        from pypdf import PdfReader
    except ImportError:
        check(False, "pypdf is installed for PDF release checks")
        return
    reader = PdfReader(pdf_path)
    metadata = reader.metadata or {}
    language = reader.trailer["/Root"].get("/Lang")
    titles = _outline_titles(reader.outline)
    check(
        metadata.get("/Title") == "Block coverage"
        and metadata.get("/Author") == "GradientCIO"
        and metadata.get("/Creator", "").startswith("Gradient CIO report renderer"),
        "PDF metadata identifies the report and renderer",
    )
    check(str(language) == "en-US", "PDF document language is en-US")
    check(
        "Block coverage" in titles and "Blocks" in titles,
        "PDF bookmarks include the cover and section headings",
    )
    fonts = _pdf_base_fonts(reader)
    require_inter = os.environ.get("REQUIRE_INTER_PDF") == "1"
    has_inter = any("inter" in name.lower() for name in fonts)
    if require_inter:
        check(has_inter, f"rendered PDF embeds or references Inter ({sorted(fonts)})")
    else:
        print(
            "  skip PDF Inter inspection"
            if not has_inter
            else "  ok   rendered PDF embeds or references Inter"
        )


def _check_libreoffice(
    check: Check, pptx_path: pathlib.Path, expected_slides: int
) -> None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    required = os.environ.get("REQUIRE_LIBREOFFICE") == "1"
    if executable is None:
        if required:
            check(False, "LibreOffice is installed for PowerPoint compatibility checks")
        else:
            print("  skip LibreOffice compatibility check (not installed)")
        return
    with tempfile.TemporaryDirectory() as directory:
        converted = pathlib.Path(directory)
        result = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(converted),
                str(pptx_path),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
        )
        pdf_path = converted / f"{pptx_path.stem}.pdf"
        pages = 0
        if result.returncode == 0 and pdf_path.is_file():
            from pypdf import PdfReader

            pages = len(PdfReader(pdf_path).pages)
        check(
            result.returncode == 0 and pages == expected_slides,
            "LibreOffice opens the golden PowerPoint without losing slides"
            + (
                f" (exit {result.returncode}, pages {pages}, expected {expected_slides})"
                if result.returncode != 0 or pages != expected_slides
                else ""
            ),
        )


def run_artifact_checks(check: Check, output_directory: pathlib.Path) -> None:
    """Render and inspect PDF/PPTX artifacts created by the release suite."""
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE

    output_directory.mkdir(parents=True, exist_ok=True)
    pptx_path = output_directory / "gradientcio-golden-report-test.pptx"
    rendered = subprocess.run(
        [
            sys.executable,
            str(GOLDEN_GENERATOR),
            "--format",
            "pptx",
            "--out",
            str(output_directory),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    check(
        rendered.returncode == 0 and pptx_path.is_file(),
        "golden generator uses the unified renderer to create PowerPoint"
        + (
            f": {(rendered.stderr or rendered.stdout).strip()[-300:]}"
            if rendered.returncode != 0
            else ""
        ),
    )
    if rendered.returncode != 0 or not pptx_path.is_file():
        return
    with tempfile.TemporaryDirectory() as second_directory:
        repeated = subprocess.run(
            [
                sys.executable,
                str(GOLDEN_GENERATOR),
                "--format",
                "pptx",
                "--out",
                second_directory,
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        repeated_path = (
            pathlib.Path(second_directory) / "gradientcio-golden-report-test.pptx"
        )
        first_digest = hashlib.sha256(pptx_path.read_bytes()).hexdigest()
        second_digest = (
            hashlib.sha256(repeated_path.read_bytes()).hexdigest()
            if repeated_path.is_file()
            else ""
        )
    check(
        repeated.returncode == 0 and first_digest == second_digest,
        "golden PowerPoint generation is byte deterministic",
    )

    presentation = Presentation(pptx_path)
    slide_text = [_slide_text(slide) for slide in presentation.slides]
    check(
        len(presentation.slides) >= 10
        and presentation.slide_width / presentation.slide_height > 1.7
        and all(text.strip() for text in slide_text),
        "PowerPoint has 16:9 structure with no empty slides",
    )

    charts = [
        shape.chart
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_chart", False)
    ]
    chart_values = [
        float(value)
        for chart in charts
        for series in chart.series
        for value in series.values
        if value is not None
    ]
    check(
        len(charts) >= 4
        and any(value < 0 for value in chart_values)
        and any(value > 0 for value in chart_values),
        "PowerPoint charts retain editable native positive and negative data",
    )
    check(
        any(chart.chart_type == XL_CHART_TYPE.DOUGHNUT for chart in charts),
        "PowerPoint renders pie blocks as editable native doughnut charts",
    )
    pacing_slides = [
        slide
        for slide, text in zip(presentation.slides, slide_text)
        if "fictional pacing metrics" in text.lower()
    ]
    check(
        len(pacing_slides) == 1
        and any(getattr(shape, "has_table", False) for shape in pacing_slides[0].shapes)
        and not any(getattr(shape, "has_chart", False) for shape in pacing_slides[0].shapes),
        "mixed-unit pacing metrics render as a table",
    )

    long_table_slides = [
        slide
        for slide, text in zip(presentation.slides, slide_text)
        if "Long table behavior" in text
    ]
    headers = []
    for slide in long_table_slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                headers.append([cell.text for cell in shape.table.rows[0].cells])
    check(
        len(long_table_slides) == 3
        and headers == [["Holding", "Weight", "Status"]] * 3
        and all("Long table behavior" in _slide_text(slide) for slide in long_table_slides),
        "split PowerPoint tables repeat headers and section headings",
    )

    all_text = "\n".join(slide_text)
    notes = "\n".join(
        slide.notes_slide.notes_text_frame.text
        for slide in presentation.slides
        if slide.has_notes_slide
    )
    check(
        "Sources: S1" in all_text
        and "Sources: S2" in all_text
        and "speaker-note retention [S2]" in notes,
        "PowerPoint retains source footers and full analysis notes",
    )

    with zipfile.ZipFile(pptx_path) as archive:
        presentation_xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.startswith("ppt/") and name.endswith(".xml")
        )
        chart_parts = [
            name
            for name in archive.namelist()
            if name.startswith("ppt/charts/chart") and name.endswith(".xml")
        ]
    check(
        "Inter" in presentation_xml and len(chart_parts) == len(charts),
        "PowerPoint stores Inter text and native chart XML",
    )

    with tempfile.TemporaryDirectory() as directory:
        invalid = pathlib.Path(directory) / "invalid.json"
        invalid.write_text(
            json.dumps({"meta": {}, "sections": [{"title": "", "blocks": []}]}),
            encoding="utf-8",
        )
        invalid_result = subprocess.run(
            [
                sys.executable,
                str(UNIFIED_RENDERER),
                str(invalid),
                "--format",
                "pptx",
                "--out",
                str(pathlib.Path(directory) / "invalid"),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        missing_result = subprocess.run(
            [
                sys.executable,
                str(UNIFIED_RENDERER),
                str(pathlib.Path(directory) / "missing.json"),
                "--format",
                "pptx",
                "--out",
                str(pathlib.Path(directory) / "missing"),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
    check(
        invalid_result.returncode == 3 and "render validation" in invalid_result.stderr,
        "unified renderer returns exit 3 for invalid report content",
    )
    check(
        missing_result.returncode == 2 and "not found" in missing_result.stderr,
        "unified renderer returns exit 2 for a missing source",
    )

    pdf_path = output_directory / "blocks.pdf"
    if pdf_path.is_file():
        _check_pdf_quality(check, pdf_path)
    else:
        check(False, "blocks PDF is available for metadata quality checks")
    _check_libreoffice(check, pptx_path, len(presentation.slides))
