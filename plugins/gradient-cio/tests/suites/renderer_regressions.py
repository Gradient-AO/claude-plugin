"""Chart, PDF, PPTX, layout, and artifact regression checks."""

from suites.harness import *
from suites.harness import _renderer_module
import check_layout as _check_layout_module
from report_layout import (
    LAYOUT_GROUP_EVIDENCE_ANALYSIS,
    LAYOUT_GROUP_TITLE_ROWS,
    LAYOUT_GROUPS_FIELD,
    LayoutPolicyError,
    build_post_render_layout_contract,
    lint_report_design,
    prepare_report_for_render,
    repair_preferred_breaks,
    resolve_layout_budget,
    validate_layout_budget,
)
from report_theme import COVER_FIELD_LIMITS

def chart_renderer():
    percentage_column = {
        "key": "ratio",
        "title": "Ratio",
        "format": "percentage",
        "decimals": 1,
    }
    check(
        _renderer_module._chart_plot_value(percentage_column, 0.25) == 25,
        "chart renderer scales percentage axes to display units",
    )
    percentage_line = _renderer_module.b_chart({"chart": {
        "title": "Percentage line",
        "status": "ok",
        "columns": [
            {"key": "period", "title": "Period", "format": "date", "decimals": None},
            percentage_column,
        ],
        "rows": [["2026-01-31", 0.25], ["2026-02-28", 0.4]],
        "render_hint": {"block": "line", "x": "period", "y": ["ratio"]},
    }})
    check(
        "<svg" in percentage_line and "%" in percentage_line,
        "chart renderer keeps same-unit percentage series as a line",
    )
    category_line = _renderer_module.b_chart({"chart": {
        "title": "Category line",
        "status": "ok",
        "columns": [
            {"key": "period", "title": "Period", "format": "text", "decimals": None},
            {"key": "value", "title": "Value", "format": "number", "decimals": 1},
        ],
        "rows": [["2026", 1.0], ["Base case", 1.5], ["2027", 2.0]],
        "render_hint": {"block": "line", "x": "period", "y": ["value"]},
    }})
    check(
        _renderer_module._line_x_kind([1, 2.5]) == "numeric"
        and _renderer_module._line_x_kind(["2026-01-31", "2026-02-28"]) == "date"
        and _renderer_module._line_x_kind(["2026", "2027"]) == "category"
        and "2026" in category_line
        and "Base case" in category_line
        and "<svg" in category_line,
        "line charts distinguish numeric, strict ISO date, and category axes",
    )
    mixed_line = _renderer_module.b_chart({"chart": {
        "title": "Mixed units",
        "status": "ok",
        "currency": "USD",
        "columns": [
            {"key": "period", "title": "Period", "format": "date", "decimals": None},
            percentage_column,
            {"key": "nav", "title": "NAV", "format": "currency", "decimals": 0},
        ],
        "rows": [["2026-01-31", 0.25, 1_000_000]],
        "truncated": False,
        "render_hint": {
            "block": "line",
            "x": "period",
            "y": ["ratio", "nav"],
        },
    }})
    check(
        "<table" in mixed_line and "<svg" not in mixed_line,
        "chart renderer falls back to a table for mixed-unit lines",
    )
    pie_chart = _renderer_module.b_chart({"chart": {
        "title": "Allocation pie",
        "status": "ok",
        "columns": [
            {"key": "sleeve", "title": "Sleeve", "format": "text", "decimals": None},
            percentage_column,
        ],
        "rows": [["Public", 0.6], ["Private", 0.3], ["Liquidity", 0.1]],
        "render_hint": {"block": "pie", "x": "sleeve", "y": ["ratio"]},
    }})
    check(
        "<svg" in pie_chart
        and "<path" in pie_chart
        and "60.0%" in pie_chart
        and _renderer_module.BLOCKS.get("pie") is _renderer_module.b_pie,
        "chart renderer supports pie hints and direct pie blocks",
    )
    mixed_pie = _renderer_module.b_chart({"chart": {
        "title": "Mixed-unit pie",
        "status": "ok",
        "columns": [
            {"key": "sleeve", "title": "Sleeve", "format": "text", "decimals": None},
            {"key": "value", "title": "Value", "format": "number", "decimals": 0},
            {"key": "unit", "title": "Unit", "format": "text", "decimals": None},
        ],
        "rows": [["Public", 60, "%"], ["Private", 30, "USD"]],
        "render_hint": {"block": "pie", "x": "sleeve", "y": ["value"]},
    }})
    check(
        "<table" in mixed_pie
        and "Public" in mixed_pie
        and "<svg" not in mixed_pie,
        "mixed-unit pie hints fall back to a table",
    )
    allocation_donut = _renderer_module.b_chart({"chart": {
        "chart_id": "bar-portfolio-hierarchy",
        "title": "Portfolio Weights by Hierarchy",
        "status": "ok",
        "columns": [
            {"key": "name", "title": "Allocation", "format": "text", "decimals": None},
            {"key": "weight", "title": "Weight", "format": "percentage", "decimals": 1},
        ],
        "rows": [["Growth", 0.55], ["Defensive", 0.35], ["Liquidity", 0.10]],
        "render_hint": {"block": "pie", "x": "name", "y": ["weight"]},
    }})
    check(
        "<svg" in allocation_donut
        and "<path" in allocation_donut
        and "A95,95" in allocation_donut
        and "55.0%" in allocation_donut
        and "<table" not in allocation_donut,
        "allocation hierarchy pie hint renders as a donut",
    )
    pacing_table = _renderer_module.b_chart({"chart": {
        "chart_id": "commitments-pacing-metrics",
        "title": "Pacing Metrics",
        "status": "ok",
        "currency": "USD",
        "columns": [
            {"key": "metric", "title": "Metric", "format": "text", "decimals": None},
            {"key": "value", "title": "Value", "format": "number", "decimals": 1},
            {"key": "unit", "title": "Unit", "format": "text", "decimals": None},
        ],
        "rows": [["Pacing gap", 0.12, "percentage"], ["Runway", 2.4, "years"]],
        "truncated": False,
        "render_hint": {"block": "table", "x": None, "y": []},
    }})
    check(
        "<table" in pacing_table and "Pacing gap" in pacing_table
        and 'class="chart donut"' not in pacing_table,
        "mixed-unit commitment pacing metrics remain a table",
    )
    categorical_line = _renderer_module.b_chart({"chart": {
        "chart_id": "categorical-line-test",
        "title": "Quarterly signal",
        "status": "ok",
        "columns": [
            {"key": "quarter", "title": "Quarter", "format": "text", "decimals": None},
            {"key": "score", "title": "Score", "format": "number", "decimals": 1},
        ],
        "rows": [["2026 Q1", 1.2], ["2026 Q2", 1.5], ["2026 Q3", 1.1]],
        "render_hint": {"block": "line", "x": "quarter", "y": ["score"]},
    }})
    check(
        "<svg" in categorical_line
        and "<path" in categorical_line
        and all(quarter in categorical_line for quarter in ("2026 Q1", "2026 Q2", "2026 Q3"))
        and "<table" not in categorical_line,
        "categorical-x chart hints render as labelled lines",
    )
    signed_bars = _renderer_module.b_chart({"chart": {
        "chart_id": "signed-bars-test",
        "title": "Attribution effects",
        "status": "ok",
        "columns": [
            {"key": "effect", "title": "Effect", "format": "text", "decimals": None},
            {"key": "return", "title": "Return", "format": "percentage", "decimals": 1},
        ],
        "rows": [["Selection", 0.004], ["Allocation", -0.006]],
        "render_hint": {"block": "bars", "x": "effect", "y": ["return"]},
    }})
    check(
        "<svg" in signed_bars
        and "0.4%" in signed_bars
        and "-0.6%" in signed_bars
        and _renderer_module.LIME in signed_bars
        and _renderer_module.CORAL in signed_bars
        and "<table" not in signed_bars,
        "negative chart hints render as signed bars",
    )
    forty_row_table = _renderer_module.b_chart({"chart": {
        "chart_id": "forty-row-table-test",
        "title": "Complete returned rows",
        "status": "ok",
        "columns": [
            {"key": "holding", "title": "Holding", "format": "text", "decimals": None},
            {"key": "weight", "title": "Weight", "format": "percentage", "decimals": 1},
        ],
        "rows": [[f"Fictional Holding {index:02d} (TEST)", index / 1000]
                 for index in range(1, 41)],
        "truncated": False,
        "render_hint": {"block": "table", "x": None, "y": []},
    }})
    check(
        forty_row_table.count("<tr>") == 41
        and "Fictional Holding 01 (TEST)" in forty_row_table
        and "Fictional Holding 40 (TEST)" in forty_row_table,
        "chart tables preserve all 40 returned rows",
    )
    coverage_html = _renderer_module.b_coverage({
        "items": [{
            "name": "Long evidence domain",
            "status": "not_assessed",
            "note": "A long note wraps independently from the complete status chip.",
        }]
    })
    coverage_document = _renderer_module.body_html({
        "meta": {
            "title": "Coverage regression (TEST)",
            "header_label": "Regression",
            "data_as_of": "2026-10-11",
        },
        "sections": [{"title": "Coverage", "blocks": [{
            "type": "coverage",
            "items": [{
                "name": "Long evidence domain",
                "status": "not_assessed",
                "note": "A long note wraps independently from the complete status chip.",
            }],
        }]}],
    })
    check(
        "not assessed" in coverage_html
        and "grid-template-columns:minmax(0,32%) max-content minmax(0,1fr)"
        in coverage_document
        and "overflow-wrap:anywhere" in coverage_document,
        "coverage rows reserve complete chip width while notes wrap",
    )
    cover_reference_document = _renderer_module.cover_html({
        "meta": {
            "title": "Cover source regression (TEST)",
            "subtitle": "Evidence summary [S1] <escaped>",
            "cover_facts": [["Evidence", "Primary record [S2]"]],
        },
    })
    deck_reference_document = _renderer_module.deck_html({
        "meta": {
            "title": "Deck source regression (TEST)",
            "subtitle": "Evidence summary [S3]",
            "cover_facts": [["Evidence", "Primary record [S4]"]],
        },
        "slides": [],
    })
    check(
        '<span class="src">S1</span>' in cover_reference_document
        and '<span class="src">S2</span>' in cover_reference_document
        and "&lt;escaped&gt;" in cover_reference_document
        and '<span class="src">S3</span>' in deck_reference_document
        and '<span class="src">S4</span>' in deck_reference_document
        and (
            f".cvr .src,.slide.dark .src{{color:{_renderer_module.INK};"
            f"background:{_renderer_module.LIME}"
        )
        in cover_reference_document,
        "PDF and deck covers render escaped source references as high-contrast chips",
    )
    compact_layout_document = _renderer_module.body_html({
        "meta": {
            "title": "Compact layout regression (TEST)",
            "header_label": "Regression",
            "data_as_of": "2026-10-11",
        },
        "sections": [{
            "title": "Compact blocks",
            "blocks": [
                {
                    "type": "tiles",
                    "tiles": [{
                        "label": "Evidence",
                        "value": "Complete",
                        "sub": (
                            "A deliberately long evidence basis wraps cleanly "
                            "beside its source [S5]"
                        ),
                    }],
                },
                {
                    "type": "two_col",
                    "left": [{
                        "type": "coverage",
                        "items": [{
                            "name": "Long evidence domain",
                            "status": "not_assessed",
                            "note": "Readable coverage detail [S6]",
                        }],
                    }],
                    "right": [{
                        "type": "statement",
                        "text": "Prominent compact conclusion [S7]",
                    }],
                },
            ],
        }],
    })
    check(
        "-webkit-line-clamp:2" in compact_layout_document
        and ".tile .ts .src{font-size:6.2pt;line-height:1.25;"
        in compact_layout_document
        and '<span class="src">S5</span>' in compact_layout_document,
        "tile subtext is limited to two lines with source-chip-safe metrics",
    )
    check(
        (
            ".two .cov{grid-template-columns:minmax(0,1fr) max-content;"
            "column-gap:8px;row-gap:2px"
        )
        in compact_layout_document
        and ".two .cov .cnote{grid-column:1/-1;font-size:8pt;"
        in compact_layout_document
        and '<span class="src">S6</span>' in compact_layout_document,
        "two-column coverage reflows notes at a readable scoped font size",
    )
    check(
        ".two .stm-t{font-size:14pt;line-height:1.25;font-weight:750}"
        in compact_layout_document
        and '<div class="stm-t">Prominent compact conclusion '
        '<span class="src">S7</span></div>' in compact_layout_document,
        "two-column statements retain prominent conclusion typography",
    )
    source_table = _renderer_module.b_table({
        "columns": ["Tag", "Evidence", "Tool"],
        "rows": [[f"S{index}", f"Evidence {index}", "test_tool"] for index in range(1, 7)],
    })
    check(
        source_table.count("<tbody>") == 1
        and "source-tail" not in source_table
        and source_table.count("<tr>") == 7
        and all(f"S{index}" in source_table for index in range(1, 7)),
        "source tables use one pageable body without a forced short tail",
    )
    table_pagination_results = []
    for row_count in (3, 4, 5, 12, 13):
        table_document = _renderer_module.body_html({
            "meta": {
                "title": f"{row_count}-row table regression (TEST)",
                "header_label": "Regression",
                "data_as_of": "2026-10-12",
            },
            "sections": [{
                "title": "Table pagination",
                "blocks": [{
                    "type": "table",
                    "title": f"Rows-{row_count}-block",
                    "columns": ["Item", "Value"],
                    "rows": [
                        [f"Fictional row {index:02d} (TEST)", index]
                        for index in range(1, row_count + 1)
                    ],
                }],
            }],
        })
        expected_class = (
            "blk table-block keep atomic-table layout-title-rows"
            if row_count <= 12
            else "blk table-block paginated-table layout-title-rows"
        )
        table_pagination_results.append(
            f'class="{expected_class}"' in table_document
            and table_document.count("<table") == 1
            and table_document.count("<thead>") == 1
            and table_document.count("<tbody>") == 1
            and table_document.count("<tr>") == row_count + 1
            and table_document.count(f"Rows-{row_count}-block") == 1
        )
    check(
        all(table_pagination_results),
        "3, 4, 5, and 12-row tables stay atomic while 13 rows use one native pageable table",
    )
    fallback_report = json.loads(
        (FIX / "attribution-compare-fallback.json").read_text(encoding="utf-8")
    )
    fallback_document = _renderer_module.body_html(fallback_report)
    returned_sign_match = re.search(
        r'<div class="btitle">Returned sign by segment</div>'
        r'(<table class="grid">.*?</table>)',
        fallback_document,
        re.DOTALL,
    )
    appendix_html = fallback_document.split(
        "Appendix — Sources and Methods",
        1,
    )[-1]
    source_table_match = re.search(
        r'(<table class="grid">.*?</table>)',
        appendix_html,
        re.DOTALL,
    )
    returned_sign_table = (
        returned_sign_match.group(1) if returned_sign_match else ""
    )
    attribution_source_table = (
        source_table_match.group(1) if source_table_match else ""
    )
    check(
        returned_sign_match is not None
        and returned_sign_table.count("<thead>") == 1
        and returned_sign_table.count("<tbody>") == 1
        and returned_sign_table.count("<tr>") == 5
        and all(
            value in returned_sign_table
            for value in (
                "Historical",
                "Governed ex ante",
                "Public Equity",
                "Alternatives",
            )
        )
        and fallback_document.count("Returned sign by segment") == 1,
        "attribution fallback keeps Returned sign by segment as one semantic table",
    )
    source_rows = next(
        block["rows"]
        for section in fallback_report["sections"]
        if section.get("id") == "appendix-a"
        for block in section.get("blocks", [])
        if block.get("type") == "table"
    )
    fixture_digests = [row[-1] for row in source_rows]
    displayed_digests = [
        f"sha256:{digest.removeprefix('sha256:')[:12]}"
        for digest in fixture_digests
    ]
    check(
        source_table_match is not None
        and attribution_source_table.count("<thead>") == 1
        and attribution_source_table.count("<tbody>") == 1
        and attribution_source_table.count("<tr>") == len(source_rows) + 1
        and attribution_source_table.count('class="digest-cell"')
        == len(source_rows)
        and all(
            (
                f'<td class="digest-cell" data-full-digest="{digest}">'
                f"{display}</td>"
            )
            in attribution_source_table
            for digest, display in zip(fixture_digests, displayed_digests)
        )
        and (
            ".grid td.digest-cell"
            "{white-space:nowrap;font-variant-numeric:tabular-nums}"
        )
        in fallback_document
        and '<div class="note digest-note">Full digests — ' in fallback_document
        and all(digest in fallback_document for digest in fixture_digests)
        and (
            ".digest-note{white-space:normal;overflow-wrap:anywhere;"
            "word-break:break-all}"
        )
        in fallback_document
        and "table-layout:fixed" not in fallback_document,
        "attribution fallback bounds digest display and prints full digest notes",
    )
    markdown_digest = "".join(_renderer_module.md_to_html(
        "| Evidence | Digest |\n"
        "| --- | --- |\n"
        f"| Fictional source (TEST) | {fixture_digests[0]} |"
    ))
    check(
        (
            f'data-full-digest="{fixture_digests[0]}">'
            f"{displayed_digests[0]}</td>"
        )
        in markdown_digest
        and markdown_digest.count(fixture_digests[0]) == 1,
        "markdown digest cells preserve the same bounded lineage contract",
    )
    check(
        (
            ".layout-title-rows .btitle,.layout-title-rows thead,"
            ".layout-title-rows tbody tr:first-child{break-after:avoid}"
        )
        in fallback_document,
        "single pageable tables keep titles, headers, and the first two data rows together",
    )
    long_failed_value = (
        "failed: governed-validation-value-without-natural-breaks-"
        + "x" * 120
    )
    validation_document = _renderer_module.body_html({
        "meta": {
            "title": "Validation wrapping regression (TEST)",
            "header_label": "Regression",
            "data_as_of": "2026-10-12",
        },
        "sections": [{
            "title": "Validation",
            "blocks": [{
                "type": "table",
                "columns": ["Tag", "Validation"],
                "rows": [["S1", long_failed_value]],
            }],
        }],
    })
    check(
        '<td class="validation-cell">' + long_failed_value + "</td>"
        in validation_document
        and ".grid td.validation-cell{white-space:normal;"
        "overflow-wrap:anywhere;word-break:break-word}" in validation_document,
        "validation cells wrap long failed values",
    )
    appendix_document = _renderer_module.body_html({
        "meta": {
            "title": "Appendix regression (TEST)",
            "header_label": "Regression",
            "data_as_of": "2026-10-11",
        },
        "sections": [
            {"title": "Main", "blocks": [{"type": "markdown", "text": "Body."}]},
            {
                "title": "Appendix A — Sources and methods",
                "blocks": [{"type": "markdown", "text": "Sources."}],
            },
        ],
    })
    check(
        '<div class="sn">A</div>' in appendix_document
        and "Appendix — Sources and methods" in appendix_document
        and "07 Appendix A" not in appendix_document,
        "JSON appendices normalize to the lettered rendering convention",
    )
    appendix_pullup_report = {
        "meta": {
            "title": "Appendix pull-up regression (TEST)",
            "header_label": "Regression",
            "data_as_of": "2026-10-12",
        },
        "sections": [
            {
                "id": "closing",
                "num": "04",
                "title": "Closing",
                "blocks": [{
                    "type": "text",
                    "text": "Short honest closing content.",
                }],
            },
            {
                "id": "sources",
                "num": "A",
                "title": "Appendix A — Sources",
                "blocks": [{
                    "type": "table",
                    "columns": ["Tag", "Source"],
                    "rows": [["S1", "Fictional source (TEST)"]],
                }],
            },
        ],
    }
    appendix_pullup_html = _renderer_module.body_html(appendix_pullup_report)
    appendix_pullup_contract = build_post_render_layout_contract(
        appendix_pullup_report,
        {
            "schema_version": 2,
            "design_profile": "analytical",
            "forced_breaks": [],
            "pdf": {
                "target_pages": {"min": 1, "max": 4},
                "hard_max_pages": 4,
                "minimum_body_fill": 0.40,
                "orphan_tail_max_fill": 0.20,
            },
            "content": {
                "defaults": {},
                "sections": {
                    "closing": {
                        "role": "closing",
                        "requires_analysis": False,
                    },
                    "sources": {
                        "role": "appendix",
                        "requires_analysis": False,
                    },
                },
            },
        },
    )
    check(
        re.search(
            r'<section class="sec">.*?<div class="sn">A</div>'
            r'.*?Appendix — Sources',
            appendix_pullup_html,
            re.DOTALL,
        )
        is not None
        and '<section class="sec new-page"><div class="keep section-atomic-table">'
        '<div class="sh"><div class="sn">A</div>' not in appendix_pullup_html
        and appendix_pullup_contract["appendix_fill"][0]["eligible"] is True,
        "first eligible appendix may follow a short closing section",
    )
    long_pacing_table = _renderer_module.b_chart({"chart": {
        "chart_id": "commitments-pacing",
        "title": "Commitments pacing",
        "status": "ok",
        "columns": [
            {"key": "period", "title": "Period", "format": "text", "decimals": None},
            {"key": "value", "title": "Value", "format": "number", "decimals": 1},
            {"key": "unit", "title": "Unit", "format": "text", "decimals": None},
        ],
        "rows": [[f"Month {index:02d}", index, "USD" if index % 2 else "%"]
                 for index in range(30)],
        "render_hint": {"block": "table", "x": None, "y": []},
    }})
    check(
        "Showing 12 of 30 rows" in long_pacing_table
        and long_pacing_table.count("<tr>") == 13
        and "Month 00" in long_pacing_table
        and "Month 16" in long_pacing_table
        and "Month 29" in long_pacing_table,
        "mixed-unit fallback tables retain endpoints and representative rows",
    )
    narrow_signed = _renderer_module.b_bars({
        "title": "Narrow signed bars",
        "narrow": True,
        "items": [
            {
                "label": "Long fictional category label",
                "value": -4.0,
                "display": "-4.0%",
            },
            {"label": "Positive", "value": 2.0, "display": "+2.0%"},
        ],
    })
    negative_bar = re.search(
        rf'<rect x="([\d.]+)"[^>]+fill="{_renderer_module.CORAL}"/>',
        narrow_signed,
    )
    negative_label = re.search(
        r'<text x="([\d.]+)"[^>]+class="cv">-4\.0%</text>',
        narrow_signed,
    )
    category_label = re.search(
        r'<text x="([\d.]+)"[^>]+class="cl">Long fictional category label</text>',
        narrow_signed,
    )
    check(
        negative_bar is not None
        and negative_label is not None
        and float(negative_label.group(1)) < float(negative_bar.group(1))
        and category_label is not None
        and float(category_label.group(1))
        >= _renderer_module._svg_text_width("Long fictional category label"),
        "narrow signed bars reserve gutters and place negative labels outside endpoints",
    )
    stress = json.loads(
        (FIX / "ic-memo-stress-report.json").read_text(encoding="utf-8")
    )
    prepared = prepare_report_for_render(stress)
    body_chart = prepared["sections"][1]["blocks"][0]["chart"]
    chart_appendix = next(
        section
        for section in prepared["sections"]
        if section.get("id") == "chart-data-appendix"
    )
    full_chart = chart_appendix["blocks"][0]["chart"]
    check(
        len(body_chart["rows"]) == 12
        and body_chart["_preview_total"] == 40
        and body_chart["_appendix_reference"] == "Appendix D — Chart data"
        and body_chart["truncated"] is False
        and len(full_chart["rows"]) == 40
        and full_chart["truncated"] is True,
        "long chart tables use a body preview and preserve full returned rows in an appendix",
    )
    stress_html = _renderer_module.body_html(stress)
    check(
        "Showing 12 of 40 returned rows; all returned rows are in Appendix D — Chart data"
        in stress_html
        and "Connector output was truncated at source" in stress_html
        and "not the complete dataset" in stress_html,
        "chart table previews link to the appendix and retain source truncation disclosure",
    )
    check(
        stress == json.loads(
            (FIX / "ic-memo-stress-report.json").read_text(encoding="utf-8")
        )
        and prepare_report_for_render(prepared) == prepared,
        "shared report preparation is source-immutable and idempotent",
    )
    layout_source = json.loads(
        (FIX / "layout-groups.json").read_text(encoding="utf-8")
    )
    layout_source["sections"][0]["blocks"][0][LAYOUT_GROUPS_FIELD] = [
        {"id": "authored:test", "kind": "custom"}
    ]
    layout_source_before = json.loads(json.dumps(layout_source))
    grouped = prepare_report_for_render(layout_source)
    grouped_table, grouped_chart, grouped_analysis = grouped["sections"][0]["blocks"]
    long_table, long_analysis = grouped["sections"][1]["blocks"]
    table_groups = grouped_table[LAYOUT_GROUPS_FIELD]
    chart_groups = grouped_chart[LAYOUT_GROUPS_FIELD]
    analysis_groups = grouped_analysis[LAYOUT_GROUPS_FIELD]
    long_groups = long_table[LAYOUT_GROUPS_FIELD]
    evidence_ids = {
        group["id"]
        for group in chart_groups
        if group.get("kind") == LAYOUT_GROUP_EVIDENCE_ANALYSIS
    }
    analysis_ids = {
        group["id"]
        for group in analysis_groups
        if group.get("kind") == LAYOUT_GROUP_EVIDENCE_ANALYSIS
    }
    check(
        any(
            group.get("kind") == LAYOUT_GROUP_TITLE_ROWS
            and group.get("minimum_rows") == 2
            for group in table_groups
        )
        and any(
            group.get("kind") == LAYOUT_GROUP_TITLE_ROWS
            and group.get("minimum_rows") == 2
            for group in chart_groups
        )
        and LAYOUT_GROUPS_FIELD
        not in grouped["sections"][2]["blocks"][0]
        and {"id": "authored:test", "kind": "custom"} in table_groups,
        "preparation groups table and nested chart titles only with two rows while preserving authored annotations",
    )
    check(
        len(evidence_ids & analysis_ids) == 1
        and not any(
            group.get("kind") == LAYOUT_GROUP_EVIDENCE_ANALYSIS
            for group in long_groups
        )
        and LAYOUT_GROUPS_FIELD not in long_analysis,
        "bounded final evidence and Analysis callouts share a group while oversized tables remain splittable",
    )
    check(
        layout_source == layout_source_before
        and prepare_report_for_render(grouped) == grouped,
        "layout grouping remains source-immutable and idempotent",
    )
    grouped_html = _renderer_module.body_html(grouped)
    check(
        "layout-title-rows" in grouped_html
        and "layout-evidence-analysis" in grouped_html
        and grouped_html.index("Titled assumptions (TEST)")
        < grouped_html.index("Alpha")
        < grouped_html.index("Beta")
        and grouped_html.index("Titled evidence chart (TEST)")
        < grouped_html.index("Analysis — Compact conclusion"),
        "PDF renderer consumes title-row and bounded evidence-analysis layout groups",
    )
    canonical_pdf = {
        "target_pages": {"min": 1, "max": 12},
        "hard_max_pages": 12,
        "minimum_body_fill": 0.40,
        "orphan_tail_max_fill": 0.20,
    }
    layout_budget = {
        "schema_version": 2,
        "design_profile": "analytical",
        "forced_breaks": ["01"],
        "pdf": canonical_pdf,
        "content": {
            "defaults": {},
            "sections": {
                "evidence": {
                    "role": "reference",
                    "requires_analysis": False,
                },
                "unbounded": {
                    "role": "reference",
                    "requires_analysis": False,
                },
                "sources": {
                    "role": "appendix",
                    "requires_analysis": False,
                },
            },
        },
    }
    check(
        validate_layout_budget(layout_budget, grouped) == []
        and lint_report_design(grouped, layout_budget) == [],
        "design lint accepts complete generated grouping and a declared configured break",
    )
    missing_group = json.loads(json.dumps(grouped))
    missing_group["sections"][0]["blocks"][1].pop(LAYOUT_GROUPS_FIELD)
    check(
        any(
            item["code"] == "design-layout-group-missing"
            for item in lint_report_design(missing_group, layout_budget)
        ),
        "design lint rejects missing title and evidence-analysis grouping",
    )
    missing_break = json.loads(json.dumps(grouped))
    missing_break["sections"][0]["page_break"] = "auto"
    missing_break["sections"][0]["new_page"] = False
    check(
        any(
            item["code"] == "design-configured-forced-break-missing"
            for item in lint_report_design(missing_break, layout_budget)
        ),
        "design lint rejects configured forced breaks absent from report JSON",
    )
    post_render_contract = build_post_render_layout_contract(
        layout_source,
        layout_budget,
    )
    check(
        post_render_contract["schema_version"] == 1
        and post_render_contract["forced_breaks"]
        == [
            {
                "configured_identifier": "01",
                "section_index": 1,
                "section_id": "evidence",
                "section_title": "Evidence",
                "source_declares_force": True,
            }
        ]
        and any(
            candidate["section_id"] == "sources"
            and candidate["eligible"] is True
            for candidate in post_render_contract["appendix_fill"]
        ),
        "post-render contract exposes forced-break expectations and bounded appendix-fill candidates",
    )
    forced_break_contract = post_render_contract["forced_breaks"]
    forced_pass = _check_layout_module.forced_break_failures(
        [
            ["Cover"],
            [
                "GradientCIO · Renderer regression",
                "Layout grouping regression (TEST)",
                "01 Evidence",
            ],
        ],
        forced_break_contract,
        grouped["meta"],
    )
    forced_fail = _check_layout_module.forced_break_failures(
        [
            ["Cover"],
            [
                "GradientCIO · Renderer regression",
                "Prior section content",
                "01 Evidence",
            ]
        ],
        forced_break_contract,
        grouped["meta"],
    )
    check(
        forced_pass == []
        and len(forced_fail) == 1
        and "forced break not applied" in forced_fail[0]
        and _check_layout_module._diagnostic(
            forced_fail[0],
            "error",
        )["code"]
        == "forced-break-not-applied",
        "post-render forced-break failures use the stable QA code",
    )
    original_bbox_pages = _check_layout_module.bbox_pages
    original_raster_metrics = _check_layout_module.raster_body_metrics
    original_page_lines = _check_layout_module.page_lines
    try:
        _check_layout_module.bbox_pages = lambda _path: [
            (792.0, []),
            (792.0, []),
            (792.0, []),
            (792.0, []),
        ]
        _check_layout_module.raster_body_metrics = lambda _path: [
            {"page": 1, "fill": 0.05, "clusters": 1},
            {"page": 2, "fill": 0.80, "clusters": 4},
            {"page": 3, "fill": 0.08, "clusters": 1},
            {"page": 4, "fill": 0.18, "clusters": 2},
        ]
        _check_layout_module.page_lines = lambda _path: [
            ["Cover"],
            ["01 Main", "Dense body content"],
            ["Divider"],
            ["A Appendix — Sources", "Sparse final appendix"],
        ]
        appendix_failures = _check_layout_module.check_layout(
            pathlib.Path("synthetic.pdf"),
            headings=[],
            minimum_fill=0.40,
            excluded={1},
            minimum_pages=None,
            maximum_pages=None,
            underfill_excluded={3},
        )
    finally:
        _check_layout_module.bbox_pages = original_bbox_pages
        _check_layout_module.raster_body_metrics = original_raster_metrics
        _check_layout_module.page_lines = original_page_lines
    check(
        any("page 4 is 18% full" in failure for failure in appendix_failures)
        and not any(
            "page 3 is 8% full" in failure for failure in appendix_failures
        )
        and any(
            "page 3 contains an orphan tail" in failure
            for failure in appendix_failures
        ),
        "minimum body fill applies to sparse final appendices while explicit dividers remain exempt",
    )
    banner_report = {
        "meta": {
            "title": "Scope banner regression (TEST)",
            "subtitle": "Subtitle remains visible.",
            "header_label": "Renderer regression",
            "data_as_of": "2026-10-12",
            "confidentiality": "Confidential (TEST)",
            "data_scope_banner": ILLUSTRATIVE_LABEL,
        },
        "sections": [
            {
                "num": "01",
                "title": "Evidence",
                "blocks": [{"type": "text", "text": "Body evidence."}],
            }
        ],
    }
    banner_cover = _renderer_module.cover_html(
        prepare_report_for_render(banner_report)
    )
    banner_body = _renderer_module.body_html(banner_report)
    banner_header, banner_footer = _renderer_module.header_footer(
        prepare_report_for_render(banner_report)
    )
    check(
        all(
            value in banner_cover
            for value in (
                "Scope banner regression (TEST)",
                "Subtitle remains visible.",
                "Confidential (TEST)",
                "Illustrative, Gradient Maintained",
            )
        )
        and "Illustrative, Gradient Maintained" in banner_header
        and "Scope banner regression (TEST)" in banner_header
        and "Confidential (TEST)" in banner_footer
        and "Evidence" in banner_body,
        "PDF cover and running header add the data-scope banner without replacing metadata",
    )
    overlong_banner = json.loads(json.dumps(banner_report))
    overlong_banner["meta"]["data_scope_banner"] = "x" * (
        COVER_FIELD_LIMITS["data_scope_banner"] + 1
    )
    try:
        prepare_report_for_render(overlong_banner)
        banner_limit_rejected = False
    except LayoutPolicyError as error:
        banner_limit_rejected = any(
            "meta.data_scope_banner" in issue for issue in error.errors
        )
    check(
        banner_limit_rejected,
        "data-scope banner uses the canonical bounded metadata limit",
    )
    invalid_budget = {
        "schema_version": 99,
        "content": {"sections": {"missing": {"max_words": 10}}},
    }
    check(
        any(
            "schema_version" in issue
            for issue in validate_layout_budget(invalid_budget, stress)
        ),
        "layout policy rejects unsupported budget schema versions",
    )
    schema_v2_report = {
        "sections": [
            {"id": "summary", "title": "Summary", "blocks": []},
            {"num": "02", "title": "Evidence", "blocks": []},
        ]
    }
    schema_v2_budget = {
        "schema_version": 2,
        "design_profile": "analytical",
        "forced_breaks": [],
        "pdf": canonical_pdf,
        "content": {
            "defaults": {"max_words": 100},
            "sections": {
                "summary": {
                    "role": "summary",
                    "requires_analysis": False,
                    "max_words": 100,
                },
                "2": {
                    "role": "analytical",
                    "requires_analysis": True,
                    "max_visuals": 2,
                },
            },
        },
    }
    check(
        validate_layout_budget(schema_v2_budget, schema_v2_report) == [],
        "schema v2 accepts roles, analysis requirements, and stable id/num resolution",
    )
    deck_report = {
        "meta": {
            "title": "Macro deck (TEST)",
            "header_label": "Macro Briefing",
            "data_as_of": "2026-10-10",
        },
        "slides": [{"title": "The message", "blocks": []}],
    }
    deck_budget = {
        "schema_version": 2,
        "design_profile": "deck",
        "forced_breaks": [],
        "pdf": canonical_pdf,
        "content": {
            "defaults": {},
            "sections": {
                "1": {"role": "reference", "requires_analysis": False}
            },
        },
    }
    prepared_deck = prepare_report_for_render(deck_report)
    check(
        "sections" not in prepared_deck
        and validate_layout_budget(deck_budget, prepared_deck) == [],
        "deck preparation preserves slides for schema-v2 budget resolution",
    )
    malformed_v2 = json.loads(json.dumps(schema_v2_budget))
    malformed_v2["design_profile"] = "brochure"
    malformed_v2["forced_breaks"] = "02"
    malformed_v2["content"]["sections"]["2"]["role"] = "narrative"
    malformed_v2["content"]["sections"]["2"]["requires_analysis"] = "yes"
    malformed_v2["content"]["sections"]["2"]["mystery"] = 1
    malformed_errors = validate_layout_budget(malformed_v2, schema_v2_report)
    check(
        all(
            any(fragment in issue for issue in malformed_errors)
            for fragment in (
                "design_profile",
                "forced_breaks",
                ".role",
                ".requires_analysis",
                "unknown field",
            )
        ),
        "schema v2 rejects wrong types, enum values, and unknown fields",
    )
    unresolved_v2 = json.loads(json.dumps(schema_v2_budget))
    unresolved_v2["content"]["sections"]["missing"] = {
        "role": "reference",
        "requires_analysis": False,
    }
    check(
        any(
            "configured section 'missing' resolves 0 times" in issue
            for issue in validate_layout_budget(unresolved_v2, schema_v2_report)
        ),
        "schema v2 requires every configured section to resolve exactly once",
    )
    mode_budget = {
        "schema_version": 2,
        "design_profile": "analytical",
        "forced_breaks": ["summary"],
        "pdf": canonical_pdf,
        "selector": {"path": "meta.mode", "default": "brief"},
        "content": {
            "defaults": {"max_words": 100},
            "sections": {
                "summary": {
                    "role": "summary",
                    "requires_analysis": False,
                }
            },
        },
        "modes": {
            "brief": {
                "forced_breaks": ["2"],
                "content": {
                    "sections": {
                        "2": {
                            "role": "analytical",
                            "requires_analysis": True,
                        }
                    }
                }
            }
        },
    }
    resolved_mode = resolve_layout_budget(mode_budget, schema_v2_report)
    check(
        validate_layout_budget(mode_budget, schema_v2_report) == []
        and resolved_mode["design_profile"] == "analytical"
        and resolved_mode["forced_breaks"] == ["2"]
        and resolved_mode["content"]["defaults"] == {"max_words": 100}
        and set(resolved_mode["content"]["sections"]) == {"summary", "2"},
        "mode budgets inherit root design profile, breaks, defaults, and sections",
    )
    mode_break_budget = json.loads(json.dumps(schema_v2_budget))
    mode_break_budget["forced_breaks"] = {
        "single": ["summary"],
        "compare": ["2"],
        "default": [],
    }
    single_mode_report = {
        **schema_v2_report,
        "meta": {"mode": "single"},
    }
    compare_mode_report = {
        **schema_v2_report,
        "meta": {"mode": "compare"},
    }
    default_mode_report = {
        **schema_v2_report,
        "meta": {"mode": "fallback"},
    }
    check(
        validate_layout_budget(mode_break_budget, single_mode_report) == []
        and resolve_layout_budget(
            mode_break_budget,
            single_mode_report,
        )["forced_breaks"]
        == ["summary"]
        and resolve_layout_budget(
            mode_break_budget,
            compare_mode_report,
        )["forced_breaks"]
        == ["2"]
        and resolve_layout_budget(
            mode_break_budget,
            default_mode_report,
        )["forced_breaks"]
        == []
        and resolve_layout_budget(
            schema_v2_budget,
            schema_v2_report,
        )["forced_breaks"]
        == [],
        "forced breaks resolve by report meta.mode while list budgets remain compatible",
    )
    lint_budget = {
        "schema_version": 2,
        "design_profile": "analytical",
        "forced_breaks": [],
        "pdf": canonical_pdf,
        "content": {
            "defaults": {},
            "sections": {
                "analysis": {
                    "role": "analytical",
                    "requires_analysis": True,
                },
                "appendix": {
                    "role": "appendix",
                    "requires_analysis": False,
                },
            },
        },
    }
    lint_report = {
        "sections": [
            {
                "id": "analysis",
                "title": "Evidence",
                "new_page": True,
                "blocks": [
                    {"type": "table", "rows": []},
                    {"type": "kv", "rows": []},
                    {
                        "type": "two_col",
                        "left": [
                            {"type": "table", "rows": []},
                            {"type": "kv", "rows": []},
                            {
                                "type": "chart",
                                "chart": {
                                    "status": "ok",
                                    "rows": [],
                                    "render_hint": {"block": "table"},
                                },
                            },
                        ],
                        "right": [],
                    },
                    {
                        "type": "callout",
                        "title": "Analysis — Analysis — Repeated",
                        "text": "Duplicate.",
                    },
                    {
                        "type": "callout",
                        "title": "Analysis — Second",
                        "text": "Second.",
                    },
                    {
                        "type": "callout",
                        "title": "Analysis — Third",
                        "text": "Third.",
                    },
                    {
                        "type": "callout",
                        "title": "Analysis — Fourth",
                        "text": "Fourth.",
                    },
                    {"type": "callout", "title": "Visual evidence", "text": ""},
                ],
            },
            {
                "id": "appendix",
                "title": "Appendix A — Sources",
                "page_break": "force",
                "blocks": [],
            },
            {
                "id": "chart-data-appendix",
                "title": "Appendix B — Chart data",
                "page_break": "force",
                "blocks": [],
            },
        ]
    }
    lint_codes = {
        item["code"] for item in lint_report_design(lint_report, lint_budget)
    }
    check(
        validate_layout_budget(lint_budget, lint_report) == []
        and {
            "design-forced-break-not-allowlisted",
            "design-analytical-evidence-missing",
            "design-analysis-callout-count",
            "design-consecutive-data-blocks",
            "design-duplicate-analysis-prefix",
            "design-generic-visual-evidence",
        }
        <= lint_codes,
        "design lint reports every analytical negative case with stable codes",
    )
    check(
        sum(
            item["code"] == "design-forced-break-not-allowlisted"
            for item in lint_report_design(lint_report, lint_budget)
        )
        == 2,
        "authored appendices require allowlisting while generated chart data is exempt",
    )
    valid_design_report = {
        "sections": [
            {
                "id": "analysis",
                "title": "Evidence",
                "blocks": [
                    {
                        "type": "callout",
                        "title": "Not available — governed result",
                        "text": "Connector response unavailable.",
                    },
                    {"type": "table", "rows": []},
                    {
                        "type": "callout",
                        "title": "Analysis — Evidence gap",
                        "text": "The gap bounds the conclusion.",
                    },
                ],
            },
            {
                "id": "appendix",
                "title": "Appendix A — Sources",
                "blocks": [],
            },
        ]
    }
    check(
        lint_report_design(valid_design_report, lint_budget) == [],
        "design lint accepts typed unavailable evidence and one analysis callout",
    )
    exempt_budget = json.loads(json.dumps(lint_budget))
    exempt_budget["design_profile"] = "reference"
    check(
        lint_report_design(lint_report, exempt_budget) == []
        and lint_report_design(
            lint_report,
            {**exempt_budget, "design_profile": "deck"},
        )
        == [],
        "reference and deck profiles are exempt from analytical design lint",
    )
    forced = prepare_report_for_render(stress)
    check(
        forced["sections"][0]["page_break"] == "force"
        and next(
            section
            for section in forced["sections"]
            if section.get("num") == "13"
        )["page_break"]
        == "prefer"
        and next(
            section
            for section in forced["sections"]
            if section.get("num") == "A"
        )["page_break"]
        == "auto",
        "layout policy preserves authored force and preferred breaks without forcing appendices",
    )
    repaired, removed = repair_preferred_breaks(
        forced,
        {"Forward Context", "Decision Analysis", "Appendix A — Sources"},
    )
    check(
        removed == ["Decision Analysis"]
        and repaired["sections"][0]["page_break"] == "force"
        and next(
            section
            for section in repaired["sections"]
            if section.get("num") == "A"
        )["page_break"]
        == "auto",
        "layout repair removes only preferred breaks and preserves automatic appendices",
    )
    cashflow_line = _renderer_module.b_chart(
        {"chart": stress["sections"][1]["blocks"][1]["chart"]}
    )
    label_positions = sorted(
        float(value)
        for value in re.findall(
            r'<text x="[^"]+" y="([\d.]+)" class="cl"[^>]*>'
            r"(?:Capital calls|Distributions|NAV|Unfunded)</text>",
            cashflow_line,
        )
    )
    check(
        "$" in cashflow_line
        and "M" in cashflow_line
        and "005437.4" not in cashflow_line
        and len(label_positions) == 4
        and all(
            right - left >= 10.9
            for left, right in zip(label_positions, label_positions[1:])
        ),
        "cash-flow axes use compact currency labels and series end labels do not overlap",
    )
    pptx_path = ROOT / "shared" / "gradient_pptx.py"
    pptx_spec = importlib.util.spec_from_file_location(
        "gradient_pptx_renderer_test", pptx_path
    )
    if pptx_spec is None or pptx_spec.loader is None:
        check(False, "PowerPoint renderer module can be loaded")
    else:
        pptx_module = importlib.util.module_from_spec(pptx_spec)
        pptx_spec.loader.exec_module(pptx_module)
        check(
            pptx_module._axis_number_format(
                5_437_400,
                decimals=1,
                currency="USD",
                value_format="currency",
            )
            == '$0.0,,"M"',
            "PowerPoint currency axes use the same compact scale",
        )
        with tempfile.TemporaryDirectory() as directory:
            pptx_output = pathlib.Path(directory) / "scope-banner.pptx"
            pptx_module.render_report(banner_report, pptx_output)
            from pptx import Presentation

            presentation = Presentation(pptx_output)
            slide_text = [
                "\n".join(
                    shape.text
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False)
                )
                for slide in presentation.slides
            ]
        check(
            len(slide_text) == 2
            and all(ILLUSTRATIVE_LABEL in text for text in slide_text)
            and all(
                value in slide_text[0]
                for value in (
                    "Scope banner regression (TEST)",
                    "Subtitle remains visible.",
                    "Confidential (TEST)",
                )
            )
            and "Evidence" in slide_text[1]
            and "Confidential (TEST)" in slide_text[1],
            "PowerPoint cover and running header add the data-scope banner without replacing metadata",
        )
        with tempfile.TemporaryDirectory() as directory:
            attribution_pptx = pathlib.Path(directory) / "attribution-fallback.pptx"
            pptx_module.render_report(fallback_report, attribution_pptx)
            attribution_presentation = Presentation(attribution_pptx)
            from pptx.enum.chart import XL_CHART_TYPE
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            attribution_charts = [
                shape.chart
                for slide in attribution_presentation.slides
                for shape in slide.shapes
                if getattr(shape, "has_chart", False)
            ]
            heat_slides = [
                slide
                for slide in attribution_presentation.slides
                if "returned cross-lane effect matrix"
                in "\n".join(
                    shape.text
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False)
                ).casefold()
            ]
            heat_tables = [
                shape.table
                for slide in heat_slides
                for shape in slide.shapes
                if getattr(shape, "has_table", False)
            ]
            heat_cells = {
                cell.text
                for table in heat_tables
                for row in table.rows
                for cell in row.cells
            }
            heat_has_picture = any(
                shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                for slide in heat_slides
                for shape in slide.shapes
            )
            with zipfile.ZipFile(attribution_pptx) as archive:
                attribution_chart_xml = "\n".join(
                    archive.read(name).decode("utf-8", errors="ignore")
                    for name in archive.namelist()
                    if name.startswith("ppt/charts/chart")
                    and name.endswith(".xml")
                )
                attribution_slide_xml = "\n".join(
                    archive.read(name).decode("utf-8", errors="ignore")
                    for name in archive.namelist()
                    if name.startswith("ppt/slides/slide")
                    and name.endswith(".xml")
                )
        check(
            sum(
                chart.chart_type == XL_CHART_TYPE.COLUMN_STACKED
                for chart in attribution_charts
            )
            >= 2
            and sum(
                chart.chart_type == XL_CHART_TYPE.LINE_MARKERS
                for chart in attribution_charts
            )
            >= 2
            and "<c:barChart" in attribution_chart_xml
            and '<c:grouping val="stacked"' in attribution_chart_xml
            and "<c:lineChart" in attribution_chart_xml,
            "attribution fallback waterfalls and lines remain editable native chart XML",
        )
        check(
            len(heat_slides) == 1
            and bool(heat_tables)
            and {"Realized", "Expected", "Allocation", "Selection", "Interaction"}
            <= heat_cells
            and not heat_has_picture
            and "<a:tbl>" in attribution_slide_xml
            and "<p:pic>" not in "\n".join(
                slide.element.xml for slide in heat_slides
            ),
            "attribution fallback heat matrix remains native editable table content",
        )

def render(keep):
    print("Render checks")
    out = ROOT / "tests" / "out" if keep else pathlib.Path(tempfile.mkdtemp())
    out.mkdir(parents=True, exist_ok=True)
    external_pdf_tools = shutil.which("pdfinfo") and shutil.which("pdftotext")
    for name, args, lo, hi, must in CASES:
        pdf = out / f"{name}.pdf"
        argv = [str(pdf) if a == "OUT" else (str(FIX / a) if (FIX / a).exists() else a) for a in args]
        r = subprocess.run([sys.executable, str(RENDER)] + argv, capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=FIX)
        if r.returncode != 0:
            check(False, f"{name}: render exit {r.returncode}: {(r.stderr or r.stdout).strip()[-300:]}"); continue
        if external_pdf_tools:
            info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, encoding="utf-8", errors="replace").stdout
            pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
            text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, encoding="utf-8", errors="replace").stdout
            size = re.search(r"Page size:\s+([\d.]+) x ([\d.]+)", info)
            page_width = float(size.group(1)) if size else None
            page_height = float(size.group(2)) if size else None
        else:
            from pypdf import PdfReader
            reader = PdfReader(pdf)
            pages = len(reader.pages)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            page_width = float(reader.pages[0].mediabox.width) if reader.pages else None
            page_height = float(reader.pages[0].mediabox.height) if reader.pages else None
        check(lo <= pages <= hi, f"{name}: {pages} pages (expected {lo}–{hi})")
        norm = lambda t: re.sub(r"\s+", "", t).upper()   # tolerate CSS uppercase and letter-spacing
        missing = [m for m in must if norm(m) not in norm(text)]
        check(not missing, f"{name}: contains {must}" + (f" — missing {missing}" if missing else ""))
        if name == "portfolio_attribution_compare_fallback" and external_pdf_tools:
            layout_json = out / f"{name}-layout.json"
            layout_checked = subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "shared" / "check_layout.py"),
                    str(pdf),
                    "--report",
                    str(FIX / "attribution-compare-fallback.json"),
                    "--min-fill",
                    "0.40",
                    "--exclude",
                    "1",
                    "--json",
                    str(layout_json),
                ],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
            )
            layout_result = (
                json.loads(layout_json.read_text(encoding="utf-8"))
                if layout_json.is_file()
                else {}
            )
            check(
                layout_checked.returncode == 0
                and not any(
                    item.get("code") == "page-underfill"
                    for item in layout_result.get("diagnostics", [])
                ),
                "attribution fallback tables preserve the 40% body-fill floor"
                + (
                    f": {(layout_checked.stderr or layout_checked.stdout).strip()[-500:]}"
                    if layout_checked.returncode
                    else ""
                ),
            )
        if name == "deck":
            check(
                page_width is not None
                and page_height is not None
                and abs(page_width - 960) < 2
                and abs(page_height - 540) < 2,
                "deck: 16:9 page size",
            )
        if name in (
            "odd",
            "digest",
            "portfolio",
            "construction_private_markets",
            "construction_fixed_income",
            "construction_global_public_equity",
            "construction_marketable_alternatives",
            "construction_real_assets",
            "portfolio_attribution_report",
            "compare",
        ):
            check("Powered by" not in text, f"{name}: no client branding by default")

    for label, prefix, lo, hi, must in GIPS_VARIANTS:
        markdown_path = FIX / f"{prefix}.md"
        visuals_path = FIX / (
            "gips_visuals.json" if prefix == "gips" else f"{prefix}-visuals.json"
        )
        meta_path = FIX / (
            "gips_meta.json" if prefix == "gips" else f"{prefix}-meta.json"
        )
        composed_path = out / f"{prefix}-composed.json"
        pdf = out / f"{prefix}.pdf"
        composed = subprocess.run(
            [
                sys.executable,
                str(GIPS_REPORT_COMPOSER),
                str(markdown_path),
                str(visuals_path),
                str(meta_path),
                str(composed_path),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if composed.returncode != 0:
            check(
                False,
                f"GIPS {label}: compose exit {composed.returncode}: "
                f"{(composed.stderr or composed.stdout).strip()[-300:]}",
            )
            continue
        rendered = subprocess.run(
            [sys.executable, str(RENDER), str(composed_path), str(pdf)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if rendered.returncode != 0:
            check(
                False,
                f"GIPS {label}: render exit {rendered.returncode}: "
                f"{(rendered.stderr or rendered.stdout).strip()[-300:]}",
            )
            continue
        if external_pdf_tools:
            info = subprocess.run(
                ["pdfinfo", str(pdf)],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
            ).stdout
            pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
            text = subprocess.run(
                ["pdftotext", "-layout", str(pdf), "-"],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
            ).stdout
        else:
            from pypdf import PdfReader
            reader = PdfReader(pdf)
            pages = len(reader.pages)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        check(
            lo <= pages <= hi,
            f"GIPS {label}: {pages} pages (expected {lo}–{hi})",
        )
        normalized = lambda value: re.sub(r"\s+", "", value).upper()
        missing = [value for value in must if normalized(value) not in normalized(text)]
        check(
            not missing,
            f"GIPS {label}: contains {must}"
            + (f" — missing {missing}" if missing else ""),
        )
    stress_pdf = out / "ic-memo-stress.pdf"
    stress_html = out / "ic-memo-stress.html"
    stress_report = FIX / "ic-memo-stress-report.json"
    stress_rendered = subprocess.run(
        [
            sys.executable,
            str(RENDER),
            str(stress_report),
            str(stress_pdf),
            "--html",
            str(stress_html),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    check(
        stress_rendered.returncode == 0,
        "IC memo stress fixture renders"
        + (
            f": {(stress_rendered.stderr or stress_rendered.stdout).strip()[-300:]}"
            if stress_rendered.returncode else ""
        ),
    )
    if stress_rendered.returncode == 0 and external_pdf_tools:
        layout_json = out / "ic-memo-stress-layout.json"
        layout_checked = subprocess.run(
            [
                sys.executable,
                str(ROOT / "shared" / "check_layout.py"),
                str(stress_pdf),
                "--report",
                str(stress_report),
                "--budget",
                str(FIX / "ic-memo-stress-budget.json"),
                "--json",
                str(layout_json),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        structured = (
            json.loads(layout_json.read_text(encoding="utf-8"))
            if layout_json.is_file()
            else {}
        )
        check(
            layout_checked.returncode == 0
            and not any(
                item.get("code") == "page-underfill"
                for item in structured.get("diagnostics", [])
            ),
            "IC memo stress fixture satisfies strict body and appendix fill"
            + (
                f": {(layout_checked.stderr or layout_checked.stdout).strip()[-500:]}"
                if layout_checked.returncode else ""
            ),
        )
        check(
            structured.get("status") == "pass"
            and structured.get("theme_version")
            and all(
                {"start_page", "end_page", "pages_touched"} <= set(section)
                for section in structured.get("sections", [])
            ),
            "layout checker emits stable structured body and section diagnostics",
        )
        qa_base = out / "ic-memo-stress-qa"
        qa_checked = subprocess.run(
            [
                sys.executable,
                str(ROOT / "shared" / "qa_report.py"),
                str(stress_report),
                "--format",
                "pdf",
                "--out",
                str(qa_base),
                "--budget",
                str(FIX / "ic-memo-stress-budget.json"),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        qa_json = out / "ic-memo-stress-qa-qa.json"
        qa_contact = out / "ic-memo-stress-qa-pdf-contact-sheet.png"
        qa_document = (
            json.loads(qa_json.read_text(encoding="utf-8"))
            if qa_json.is_file()
            else {}
        )
        check(
            qa_checked.returncode == 0
            and qa_document.get("status") == "pass"
            and qa_document.get("outputs", {}).get("pdf", {}).get("status")
            == "pass"
            and qa_document.get("repair", {}).get("rerendered") is False
            and qa_document.get("repair", {}).get(
                "removed_preferred_breaks"
            )
            == []
            and qa_contact.is_file(),
            "one-command QA writes structured diagnostics and a labeled PDF contact sheet"
            + (
                f": {(qa_checked.stderr or qa_checked.stdout).strip()[-500:]}"
                if qa_checked.returncode
                else ""
            ),
        )
    run_artifact_checks(check, out)
    print(f"PDFs in {out}")
